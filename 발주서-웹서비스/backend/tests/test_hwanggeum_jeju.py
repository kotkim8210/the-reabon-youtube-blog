"""황금향 제주다팜 발주·토스 수집·운송장(한진택배) (2026-10-16 신규).

판매옵션 9종: '하우스 정품 황금향 {소과(100g~150g)|중과(150~220g)|중대과(220~300g)} {2|3|5}kg'.
쿠팡 '[재구매1위] 제주 황금향 정품 하우스 고당도 평균16brix 산지직송 맛보장' 옵션 '1박스 2kg (소)'(DeliveryList 2026-10-16 실측).
"""

import asyncio
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import order_routing
from app.processors import kolrabi_order, tomato_order, tracking_input
from app.processors.tracking_input import _semantic_option_keys

COUPANG_PRODUCT = "[재구매1위] 제주 황금향 정품 하우스 고당도 평균16brix 산지직송 맛보장"
SHINANO = "시나노 골드 사과 노란사과 황금사과 경북 가정용 상등급"


def test_all_9_options_map():
    for grade in ("소과", "중과", "중대과"):
        for kg in ("2", "3", "5"):
            got = kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, f"1박스 {grade} {kg}kg")
            assert got and got.startswith(f"하우스 정품 황금향 {grade}(") and got.endswith(f" {kg}kg"), (grade, kg, got)
    assert len(set(kolrabi_order._HWANGGEUM_JEJU_OPTIONS.values())) == 9


def test_real_and_short_grade_options():
    assert (kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "1박스 2kg (소)")
            == "하우스 정품 황금향 소과(100g~150g) 2kg")
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "1박스 3kg (중)") == "하우스 정품 황금향 중과(150~220g) 3kg"
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "1박스 5kg (중대)") == "하우스 정품 황금향 중대과(220~300g) 5kg"
    # 제주다팜엔 대과가 없어 가장 큰 중대과로
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "1박스 2kg (대)") == "하우스 정품 황금향 중대과(220~300g) 2kg"
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "대과 3kg") == "하우스 정품 황금향 중대과(220~300g) 3kg"
    # 등급 없음·없는 중량은 변환 실패(발주서 needs_check로 표시)
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "1박스 2kg") is None
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "1박스 10kg (소)") is None
    # '소량'·'대용량' 같은 낱말 속 한 글자는 등급이 아님
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "대용량 2kg") is None


def test_exclusive_from_shinano_and_routing():
    assert kolrabi_order.convert_hwanggeum_option(SHINANO, "1박스 소과 2kg 10~13과(한입)") is None
    assert kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 2kg (소)") is None
    assert order_routing.resolve_order_supplier(COUPANG_PRODUCT, "1박스 2kg (소)") == "제주다팜"


def _delivery(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 3, r.get("order_no", f"OID{i}"))
        ws.cell(row, 11, r["product"])
        ws.cell(row, 12, r["option"])
        ws.cell(row, 23, r.get("qty", 1))
        ws.cell(row, 27, r.get("name", f"수취인{i}"))
        ws.cell(row, 28, r.get("phone", "010-0000-0000"))
        ws.cell(row, 29, "12345")
        ws.cell(row, 30, r.get("address", "서울시 어딘가"))
        ws.cell(row, 31, "")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_process_outputs_gives_separate_hwanggeum_and_shinano_files():
    data = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1박스 2kg (소)", "name": "김광은", "order_no": "3103637528262"},
        {"product": SHINANO, "option": "1박스 소과 2kg 10~13과(한입)", "name": "김광은", "order_no": "3103637528262"},
    ])
    results = {r[2]["product"]: r for r in kolrabi_order.process_outputs(data)}
    assert set(results) == {"황금향(제주다팜)", "시나노골드사과(제주다팜)"}
    out, filename, stats = results["황금향(제주다팜)"]
    assert "황금향" in filename and stats["total"] == 1
    ws = load_workbook(BytesIO(out)).active
    assert (ws.cell(2, 2).value, ws.cell(2, 8).value) == ("김광은", "하우스 정품 황금향 소과(100g~150g) 2kg")
    assert stats["options"][0]["orders"] == [{"order_id": "3103637528262", "quantity": 1}]


class _FakeTossClient:
    def __init__(self, orders):
        self._orders = orders
        self.registered = []

    async def get_orders(self, start_date=None, end_date=None, status=None):
        return self._orders

    async def register_tracking(self, order_product_id, delivery_company, tracking_number):
        self.registered.append((order_product_id, delivery_company, tracking_number))
        return {"ok": True}


def test_toss_hwanggeum_collected(monkeypatch):
    import app.toss.client as toss_module

    orders = [{"productName": COUPANG_PRODUCT, "optionName": "1박스 3kg (중)", "quantity": 1,
               "receiverName": "토스손님", "receiverPhone": "010-1111-2222", "orderProductStatus": "PAY_DONE"}]
    monkeypatch.setattr(toss_module, "toss_client", _FakeTossClient(orders), raising=False)
    result = asyncio.run(tomato_order.collect_toss_jejudapam_orders("2026-10-15", "2026-10-16"))
    assert [(e["name"], e["product"]) for e in result["hwanggeum"]] == [
        ("토스손님", "하우스 정품 황금향 중과(150~220g) 3kg")]
    assert result["shinano"] == []


def test_semantic_keys_match_both_sides():
    ol = _semantic_option_keys("하우스 정품 황금향 소과(100g~150g) 2kg", "")
    dl = _semantic_option_keys(COUPANG_PRODUCT, "1박스 2kg (소)")
    assert "hwanggeum:소과:2kg" in ol and "hwanggeum:소과:2kg" in dl
    assert "hwanggeum:중대과:5kg" in _semantic_option_keys("하우스 정품 황금향 중대과(220~300g) 5kg")


def _orderlist(rows: list[dict]) -> bytes:
    """제주다팜 회신: G 상품명·K 수령인·M 연락처·O 주소·Q 택배사·R 운송장."""
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 7, r["product"])
        ws.cell(row, 11, r["name"])
        ws.cell(row, 13, r.get("phone", "010-0000-0000"))
        ws.cell(row, 15, r.get("address", "서울시 어딘가"))
        ws.cell(row, 17, r.get("courier", ""))
        ws.cell(row, 18, r["tracking"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracking_hwanggeum_hanjin_default_and_same_person_split():
    """같은 사람이 황금향+시나노골드 — 각 행에 맞는 송장, 황금향은 회신 택배사 빈칸이면 한진택배."""
    delivery = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1박스 2kg (소)", "name": "김광은", "order_no": "3103637528262"},
        {"product": SHINANO, "option": "1박스 소과 2kg 10~13과(한입)", "name": "김광은", "order_no": "3103637528262"},
    ])
    orderlist = _orderlist([
        {"product": "시나노골드(가정용) 소과 2kg(10-13과내외)", "name": "김광은", "tracking": "254000000001", "courier": "롯데택배"},
        {"product": "하우스 정품 황금향 소과(100g~150g) 2kg", "name": "김광은", "tracking": "512000000002"},  # 택배사 빈칸
    ])
    out, _fn, stats = tracking_input.process(orderlist, delivery)
    ws = load_workbook(BytesIO(out)).active
    assert (ws.cell(2, 4).value, ws.cell(2, 5).value) == ("한진택배", "512000000002")
    assert (ws.cell(3, 4).value, ws.cell(3, 5).value) == ("롯데택배", "254000000001")
    assert stats["filled"] == 2


def test_tracking_reply_hanjin_text_normalized():
    delivery = _delivery([{"product": COUPANG_PRODUCT, "option": "1박스 3kg (중)", "name": "홍길동"}])
    orderlist = _orderlist([{"product": "하우스 정품 황금향 중과(150~220g) 3kg", "name": "홍길동",
                             "tracking": "512000000003", "courier": "한진"}])
    out, _fn, _stats = tracking_input.process(orderlist, delivery)
    ws = load_workbook(BytesIO(out)).active
    assert (ws.cell(2, 4).value, ws.cell(2, 5).value) == ("한진택배", "512000000003")


def test_toss_tracking_hwanggeum_registers_hanjin_without_warning(monkeypatch):
    """토스 운송장: 황금향은 회신 택배사가 비어도 한진택배로 등록하고 '롯데 아님' 경고를 띄우지 않는다."""
    import app.processors.tomato_tracking as tt

    fake = _FakeTossClient([{
        "productName": COUPANG_PRODUCT, "optionName": "1박스 2kg (소)", "quantity": 1, "orderProductId": "OP1",
        "orderId": "T1", "receiverName": "토스손님", "receiverPhone": "010-1111-2222",
        "address": "서울시 어딘가", "orderProductStatus": "PREPARING_PRODUCT",
    }])
    monkeypatch.setattr(tt, "toss_client", fake, raising=False)
    monkeypatch.setattr(tt, "parse_reply_file", lambda _b: [{
        "name": "토스손님", "phone": "01011112222", "phone_digits": "01011112222", "address": "서울시어딘가",
        "tracking": "512000000009", "courier": "", "option_keys": set(),
    }])
    stats = asyncio.run(tt.process_toss_watermelon_tracking(b"x"))
    assert fake.registered == [("OP1", "한진택배", "512000000009")]
    assert stats["toss_success"] == 1 and not stats.get("toss_courier_warnings")
