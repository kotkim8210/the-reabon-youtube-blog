"""고구마 등급 '대' 판정 — 쿠팡 옵션 '2kg (대: 튀김/요리/생식용)' 누락 회귀 (2026-10-19 최광균 건).

종전 정규식은 '(대)'·'kg 대'처럼 바로 닫히는 형태만 인정해, 괄호 안 등급 뒤에 설명이 붙으면
발주서 N열에 원문이 그대로 들어갔다('호박고구마 2Kg (대)'가 아니라 '1박스 2kg (대: …) …').
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.processors import hobak_goguma
from app.processors.goguma_order import canonical_goguma_option

HOBAK = "국내산 해남 황토밭 햇 호풍미 호박고구마 판매1위"
KKUL = "국내산 해남 황토밭 황금 꿀고구마"


def test_dae_with_colon_description_is_recognised():
    assert canonical_goguma_option("1박스 2kg (대: 튀김/요리/생식용)", HOBAK) == "호박고구마 2Kg (대)"
    assert canonical_goguma_option("1박스 5kg (대 : 튀김/요리/생식용)", HOBAK) == "호박고구마 5Kg (대)"
    assert canonical_goguma_option("1박스 3kg(대)", HOBAK) == "호박고구마 3Kg (대)"
    assert canonical_goguma_option("1박스 2kg (대: 튀김/요리/생식용)", KKUL) == "꿀고구마 2Kg (대)"
    # 다른 등급은 종전대로
    assert canonical_goguma_option("1박스 2kg (한입: 에프용/아이간식)", HOBAK) == "호박고구마 2Kg (한입)"
    assert canonical_goguma_option("1박스 3kg (중상: 구이/찜 겸용 인기)", HOBAK) == "호박고구마 3Kg (중상)"
    assert canonical_goguma_option("1박스 5kg (특상: 선물용)", HOBAK) == "호박고구마 5Kg (특상)"
    # 상품명 속 '대'(판매1위·대용량·해남대…)는 등급이 아니다 — 등급 없는 옵션은 변환 실패가 맞다
    assert canonical_goguma_option("1박스 3kg", "대용량 호박고구마 해남대 판매1위") == "1박스 3kg 대용량 호박고구마 해남대 판매1위"


def test_hobak_order_row_with_dae_colon_goes_into_file_without_needs_check():
    wb = Workbook()
    ws = wb.active
    ws.append(["헤더"] * 31)
    row = [""] * 31
    row[2], row[10], row[11], row[22], row[26], row[27], row[28], row[29], row[30] = (
        "4103672625025", HOBAK, "1박스 2kg (대: 튀김/요리/생식용)", 1, "최광균", "0504-2788-2574", "03665",
        "서울특별시 서대문구 남가좌동 344-4 203호", "문 앞")
    ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    out, _fn, stats = hobak_goguma.process_order(buf.getvalue(), [])
    assert stats["total"] == 1 and "needs_check" not in stats
    ws2 = load_workbook(BytesIO(out)).active
    assert (ws2.cell(2, 1).value, ws2.cell(2, 14).value) == ("최광균", "호박고구마 2Kg (대)")
