"""라이브 이벤트 당첨자 → 햇밤은 쥬얼리프룻 발주서 (2026-10-26).

제주다팜엔 햇밤이 없다(adminplus 검색: 밤고구마·미니밤호박뿐). 경품 '햇밤1kg' → 쿠팡 주문과 같은
'햇밤 대 1kg'로 쥬얼리프룻 양식에 따로 낸다. 다른 경품과 섞이면 발주처별로 파일이 나뉜다.
"""

from io import BytesIO

from openpyxl import load_workbook

from app.processors import event_order

HEADER = "경품명,별명,개인 식별 정보 확인,이름,연락처,주소,주문 아이디,구매 금액,환불/취소 금액,환불/날짜/시간\n"


def _csv(rows: list[str]) -> bytes:
    return (HEADER + "\n".join(rows) + "\n").encode("utf-8-sig")


def test_chestnut_prize_routes_to_jewelry():
    assert event_order.event_product_name("햇밤1kg") == "햇밤 대 1kg"
    assert event_order.event_product_name("햇밤 3kg") == "햇밤 대 3kg"
    assert event_order.event_supplier("햇밤1kg") == "jewelry"
    # 밤고구마·미니밤호박은 햇밤이 아니다
    assert event_order.event_supplier("제주 미니밤호박 1kg") == "jejudapam"
    assert event_order.event_supplier("감홍사과1.5kg") == "jejudapam"


def test_process_chestnut_only_csv_gives_one_jewelry_file():
    data = _csv([
        "햇밤1kg,별명1,2026/10/24 21:01,테스트가,010-0000-0001,부산광역시 기장군 테스트로 1,7103820345139,7700,0,",
        "햇밤1kg,별명2,2026/10/24 21:01,테스트나,010-0000-0002,경기도 파주시 테스트로 2,18103820340130,96660,0,",
    ])
    results = event_order.process(data)
    assert len(results) == 1
    out, filename, stats = results[0]
    assert filename.startswith("이벤트당첨_쥬얼리프룻_햇밤_발주(") and stats["supplier"] == "쥬얼리프룻"
    assert stats["winners"] == 2
    ws = load_workbook(BytesIO(out)).active
    rows = [(ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 7).value, ws.cell(r, 8).value, ws.cell(r, 13).value)
            for r in (2, 3)]
    assert rows == [
        ("아이티소프트", "테스트가", "햇밤 대 1kg", "1", "7103820345139"),
        ("아이티소프트", "테스트나", "햇밤 대 1kg", "1", "18103820340130"),
    ]
    assert ws.cell(2, 9).value == "식품애착"


def test_mixed_csv_splits_jewelry_and_jejudapam():
    data = _csv([
        "햇밤1kg,별명1,2026/10/24 21:01,밤당첨,010-0000-0001,부산광역시 테스트로 1,7103820345139,7700,0,",
        "감홍사과1.5kg,별명2,2026/10/25 22:01,감홍당첨,010-0000-0002,대구광역시 테스트로 2,21103842199632,34400,0,",
    ])
    results = event_order.process(data)
    by_supplier = {st["supplier"]: (fn, st) for _b, fn, st in results}
    assert set(by_supplier) == {"쥬얼리프룻", "제주다팜"}
    assert by_supplier["쥬얼리프룻"][1]["winners"] == 1 and by_supplier["제주다팜"][1]["winners"] == 1
    assert "감홍사과" in by_supplier["제주다팜"][0] and "햇밤" in by_supplier["쥬얼리프룻"][0]
