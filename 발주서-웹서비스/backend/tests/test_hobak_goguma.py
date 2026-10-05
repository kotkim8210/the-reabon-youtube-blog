"""햇 호박고구마(해달) 별도 섹션 — 발주·쿠팡 운송장·토스 운송장 (2026-09-28 신규).

실측: 쿠팡 itsoft001 DeliveryList K='국내산 해남 황토밭 햇 호풍미 호박고구마 판매1위',
L='1박스 10kg (특상: 군고구마용)', D열(택배사)은 '롯데택배'로 미리 채워져 나온다.
토스: productName='국내산 해남 황토밭 햇 호풍미 호박고구마 판매1위, 2kg, 1박스, 한입'.

사고 방지 포인트
- 토스 수집이 '고구마' 키워드라 호박고구마가 꿀고구마 발주서에 '꿀고구마 2Kg (한입)'으로 섞였다.
- itsoft001 D열 '롯데택배'를 그대로 두면 해달(한진) 송장이 롯데로 등록된다.
"""

import asyncio
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.processors import goguma_order, hobak_goguma, toss_auto
from app.processors.tracking_match import coupang_courier_name

HOBAK_PRODUCT = "국내산 해남 황토밭 햇 호풍미 호박고구마 판매1위"
HOBAK_OPTION = "1박스 10kg (특상: 군고구마용)"
TOSS_HOBAK = "국내산 해남 황토밭 햇 호풍미 호박고구마 판매1위, 2kg, 1박스, 한입"
GALBI_PRODUCT = "한입 LA갈비 양념갈비 구매 1회 특급쉐프소스 양념소갈비"


# ── 품목명 ──
def test_hobak_names_are_hobak_not_kkul():
    assert goguma_order.canonical_goguma_option(HOBAK_OPTION, HOBAK_PRODUCT) == "호박고구마 10Kg (특상)"
    assert goguma_order.transform_toss_option(TOSS_HOBAK, "") == "호박고구마 2Kg (한입)"
    # 꿀고구마는 종전 그대로
    assert goguma_order.canonical_goguma_option("0. 해남 황금 꿀고구마: 5Kg 중상[추천]") == "꿀고구마 5Kg (중상)"
    assert goguma_order.transform_toss_option("해남 황금 꿀고구마 3kg 한입", "") == "꿀고구마 3Kg (한입)"


# ── 토스 수집 분리 (중복 발주 방지) ──
def test_toss_filters_split_hobak_from_kkul():
    hobak = {"productName": TOSS_HOBAK, "optionName": ""}
    kkul = {"productName": "해남 황금 꿀고구마", "optionName": "3kg 중상"}
    assert goguma_order.is_goguma_order(hobak) is False       # 꿀고구마 페이지는 호박고구마를 안 잡는다
    assert toss_auto.is_goguma_order(hobak) is False
    assert goguma_order.is_hobak_goguma_toss_order(hobak) is True
    assert goguma_order.is_goguma_order(kkul) is True
    assert toss_auto.is_goguma_order(kkul) is True
    assert goguma_order.is_hobak_goguma_toss_order(kkul) is False


def _delivery(rows: list[dict]) -> bytes:
    """쿠팡 DeliveryList (C 주문번호, D 택배사='롯데택배', K 상품명, L 옵션, W 수량, AA~AE 수취인)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["헤더"] * 41)
    for i, r in enumerate(rows):
        row = [""] * 41
        row[2] = r.get("order_no", f"OID{i}")
        row[3] = "롯데택배"
        row[10] = r["product"]
        row[11] = r["option"]
        row[22] = r.get("qty", 1)
        row[26] = r["name"]
        row[27] = r.get("phone", f"0504-0000-000{i}")
        row[28] = "13303"
        row[29] = r.get("address", f"경기도 성남시 어딘가 {i}")
        row[30] = "문 앞"
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_process_order_merges_coupang_and_toss_only_hobak():
    data = _delivery([
        {"product": GALBI_PRODUCT, "option": "800g 2개", "name": "최성옥"},
        {"product": HOBAK_PRODUCT, "option": HOBAK_OPTION, "name": "변선아", "order_no": "31103128445807"},
        {"product": "2026 햇 안동 고당도 홍로사과", "option": "1박스 소과 1.5kg(7-10과내)", "name": "김태양"},
    ])
    toss_entries = [{
        "name": "토스손님", "phone": "050215460912", "zipcode": "17816", "address": "경기도 평택시",
        "qty": "1", "product": goguma_order.transform_toss_option(TOSS_HOBAK, ""), "memo": "",
        "order_id": "288697616",
    }]
    out, filename, stats = hobak_goguma.process_order(data, toss_entries)
    assert "호박고구마" in filename and "아이티소프트" in filename
    assert stats["total"] == 2 and stats["coupang"] == 1 and stats["toss"] == 1
    assert "needs_check" not in stats
    ws = load_workbook(BytesIO(out)).active
    rows = [(ws.cell(r, 1).value, ws.cell(r, 14).value, ws.cell(r, 7).value) for r in range(2, ws.max_row + 1)
            if ws.cell(r, 1).value]
    assert rows == [("변선아", "호박고구마 10Kg (특상)", "식품애착"), ("토스손님", "호박고구마 2Kg (한입)", "식품애착")]
    # 중복발주 이력 키: 쿠팡=주문번호|L열 원문, 토스=주문번호|발주명
    from app.processors.issued_orders import order_ids_from_stats
    keys = set(order_ids_from_stats(stats))
    assert "31103128445807|1박스10kg(특상:군고구마용)" in keys
    assert "288697616|호박고구마2Kg(한입)" in keys


def test_issued_roundtrip_excludes_yesterdays_coupang_row():
    from app.processors.issued_orders import filter_delivery_by_issued, order_ids_from_stats

    data = _delivery([{"product": HOBAK_PRODUCT, "option": HOBAK_OPTION, "name": "변선아", "order_no": "31103128445807"}])
    _out, _fn, first = hobak_goguma.process_order(data, [])
    issued = set(order_ids_from_stats(first))
    names: list[str] = []
    filtered, skipped = filter_delivery_by_issued(data, issued, skipped_names=names)
    assert skipped == 1 and names == ["변선아"]
    _out2, _fn2, second = hobak_goguma.process_order(filtered, [])
    assert second["total"] == 0


def test_unparsed_option_is_flagged_not_silent():
    data = _delivery([{"product": HOBAK_PRODUCT, "option": "1박스 선물용", "name": "이상옵션"}])
    _out, _fn, stats = hobak_goguma.process_order(data, [])
    assert stats["total"] == 1
    assert any("이상옵션" in s for s in stats["needs_check"])


# ── 쿠팡 운송장 (DeliveryList E·D열) ──
_HANJIN_HEADERS = [
    "받으시는 분", "받으시는 분 전화", "받는분담당자(선택)", "받는분핸드폰(선택)",
    "받는분우편번호(선택)", "받는분총주소", "보내시는 분", "보내시는 분 전화",
    "보내는분담당자(선택)", "보내는분담당자HP(선택)", "보내는분우편번호(선택)", "보내는분총주소",
    "수량", "품목명", "운임Type", "지불조건", "출고번호", "특기사항", "메모1", "메모2", "메모3", "메모4",
]


def _haedal_reply(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(_HANJIN_HEADERS)
    for r in rows:
        row = [""] * 22
        row[0] = r["name"]
        row[1] = r.get("phone", "")
        row[5] = r.get("address", "")
        row[6] = "식품애착"
        row[12] = 1
        row[13] = r["product"]
        row[15] = r.get("pay", "선불")
        row[16] = r["tracking"]
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracking_fills_only_hobak_rows_and_forces_hanjin_courier():
    dl = _delivery([
        {"product": GALBI_PRODUCT, "option": "800g 2개", "name": "변선아", "phone": "0504-1111-2222"},   # 동명이인(다른 상품)
        {"product": HOBAK_PRODUCT, "option": HOBAK_OPTION, "name": "변선아", "phone": "0504-3333-4444"},
    ])
    reply = _haedal_reply([
        {"name": "변선아", "phone": "0504-3333-4444", "product": "호박고구마 10Kg (특상)", "tracking": "463319275999"},
    ])
    out, filename, stats = hobak_goguma.process_tracking(reply, dl)
    assert "호박고구마" in filename
    assert stats["filled"] == 1 and stats["skipped"] == 0, stats
    ws = load_workbook(BytesIO(out)).active
    # 2행(LA갈비 변선아)은 그대로, 3행(호박고구마 변선아)에만 송장·한진택배
    assert ws.cell(2, 5).value in (None, "") and ws.cell(2, 4).value == "롯데택배"
    assert ws.cell(3, 5).value == "463319275999"
    assert ws.cell(3, 4).value == "한진택배"


def test_hanjin_short_name_maps_to_coupang_label():
    assert coupang_courier_name("한진") == "한진택배"
    assert coupang_courier_name("한진택배") == "한진택배"
    assert coupang_courier_name("롯데") == "롯데택배"


# ── 토스 운송장 (API, 스텁) ──
def test_toss_tracking_registers_only_hobak_orders(monkeypatch):
    reply = _haedal_reply([
        {"name": "문유금", "phone": "0502-1546-0912", "address": "경기도 평택시 안중읍",
         "product": "호박고구마 2Kg (한입)", "tracking": "463319275871"},
    ])
    toss_items = [
        {   # 호박고구마 — 등록 대상
            "orderProductId": "OP-HOBAK", "orderId": "288697616", "productName": TOSS_HOBAK, "optionName": "",
            "receiverName": "문유금", "receiverRealPhone": "050215460912",
            "address": "경기도 평택시 안중읍", "detailAddress": "", "orderProductStatus": "PAID",
            "shippingTrackingNumber": None,
        },
        {   # 같은 사람 꿀고구마 — 호박 섹션에서는 건드리면 안 됨
            "orderProductId": "OP-KKUL", "orderId": "288697617", "productName": "해남 황금 꿀고구마",
            "optionName": "3kg 중상", "receiverName": "문유금", "receiverRealPhone": "050215460912",
            "address": "경기도 평택시 안중읍", "detailAddress": "", "orderProductStatus": "PAID",
            "shippingTrackingNumber": None,
        },
    ]

    async def fake_get_orders(**_kwargs):
        return toss_items

    registered: list[tuple] = []

    async def fake_register(order_product_id, delivery_company, tracking_number):
        registered.append((order_product_id, delivery_company, tracking_number))
        return {"ok": True}

    monkeypatch.setattr(toss_auto.toss_client, "get_orders", fake_get_orders)
    monkeypatch.setattr(toss_auto.toss_client, "register_tracking", fake_register)

    result = asyncio.run(toss_auto.process_toss_tracking(
        reply, product_filter=goguma_order.is_hobak_goguma_toss_order, product_label="호박고구마",
    ))
    assert result["success"] == 1 and result["fail"] == 0, result
    assert registered == [("OP-HOBAK", "한진택배", "463319275871")]   # 회신 택배사 빈칸 → 해달 기본 한진


# ── 라이브 이벤트 당첨자 (2026-10-06) — 실제 당첨자 개인정보 대신 가짜 데이터 사용 ──
_WINNERS_HEADER = "경품명,별명,개인 식별 정보 확인,이름,연락처,주소,주문 아이디,구매 금액,환불/취소 금액,환불/날짜/시간"


def _winners_csv(rows: list[str]) -> bytes:
    return ("\n".join([_WINNERS_HEADER, *rows]) + "\n").encode("utf-8-sig")


def test_event_winners_become_haedal_rows():
    data = _winners_csv([
        "호박고구마 2kg(중상),김*희,2026/10/05 21:01,김테스트,010-1111-2222,서울특별시 어딘가 1,27100000000001,21800,0,",
        "호박고구마 2kg(중상),태풍,2026/10/05 21:01,이테스트,010-3333-4444,경기도 어딘가 2,12100000000002,8400,0,",
    ])
    out, filename, stats = hobak_goguma.process_event(data)
    assert "이벤트당첨" in filename and "호박고구마" in filename
    assert stats["event"] == 2 and stats["total"] == 2
    assert "needs_check" not in stats
    ws = load_workbook(BytesIO(out)).active
    rows = [(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 6).value, ws.cell(r, 14).value,
             ws.cell(r, 13).value, ws.cell(r, 19).value, ws.cell(r, 7).value)
            for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    assert rows == [
        ("김테스트", "010-1111-2222", "서울특별시 어딘가 1", "호박고구마 2Kg (중상)", "1", "문 앞", "식품애착"),
        ("이테스트", "010-3333-4444", "경기도 어딘가 2", "호박고구마 2Kg (중상)", "1", "문 앞", "식품애착"),
    ]


def test_event_skips_refunds_and_flags_other_prizes():
    data = _winners_csv([
        "호박고구마 3kg(특상),a,x,정상당첨,010-1,서울 1,1001,1,0,",
        "호박고구마 2kg(중상),b,x,환불당첨,010-2,서울 2,1002,1,8400,2026/10/05 22:00",   # 환불 → 제외
        "미니밤호박 3kg,c,x,다른경품,010-3,서울 3,1003,1,0,",                              # 호박고구마 아님
    ])
    out, _fn, stats = hobak_goguma.process_event(data)
    assert stats["event"] == 1 and stats["refund_skipped"] == 1
    assert any("다른경품" in s for s in stats["needs_check"])
    ws = load_workbook(BytesIO(out)).active
    assert ws.cell(2, 1).value == "정상당첨" and ws.cell(2, 14).value == "호박고구마 3Kg (특상)"
    assert ws.cell(3, 1).value in (None, "")


def test_event_reupload_on_later_day_is_excluded():
    """같은 당첨자 CSV를 다음 날 또 올려도 경품이 두 번 나가지 않는다."""
    from app.processors.issued_orders import order_ids_from_stats

    data = _winners_csv(["호박고구마 2kg(중상),a,x,김테스트,010-1,서울 1,27100000000001,1,0,"])
    _out, _fn, first = hobak_goguma.process_event(data)
    issued = set(order_ids_from_stats(first))
    assert issued == {"27100000000001|호박고구마2Kg(중상)"}
    names: list[str] = []
    try:
        hobak_goguma.process_event(data, exclude_keys=issued, skipped_names=names)
    except ValueError as exc:
        assert "김테스트" in str(exc)
    else:
        raise AssertionError("이미 발주한 당첨자는 다시 발주서에 들어가면 안 된다")
