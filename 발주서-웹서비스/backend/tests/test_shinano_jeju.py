"""시나노골드 사과 제주다팜 발주·토스 수집·운송장 (2026-10-16 신규).

공급처 구성상품 30종: '시나노골드(가정용) {대과|중과|중소과|소과|중대과랜덤} {1.5|2|2.5|3|4|5}kg(과수)'.
쿠팡 '시나노 골드 사과 노란사과 황금사과 경북 가정용 상등급' 옵션 '1박스 소과 2kg 10~13과(한입)'(DeliveryList 2026-10-16 실측).
"""

import asyncio
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import order_routing
from app.processors import kolrabi_order, tomato_order, tracking_input
from app.processors.tracking_input import _semantic_option_keys

COUPANG_PRODUCT = "시나노 골드 사과 노란사과 황금사과 경북 가정용 상등급"
HWANGGEUM = "[재구매1위] 제주 황금향 정품 하우스 고당도 평균16brix 산지직송 맛보장"


def test_all_30_options_map_to_jeju_names():
    for grade in ("대과", "중과", "중소과", "소과", "중대과랜덤"):
        for kg in ("1.5", "2", "2.5", "3", "4", "5"):
            got = kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, f"1박스 {grade} {kg}kg")
            assert got and got.startswith(f"시나노골드(가정용) {grade} {kg}kg("), (grade, kg, got)
    assert len(set(kolrabi_order._SHINANO_JEJU_OPTIONS.values())) == 30


def test_real_coupang_option_and_grade_collisions():
    assert (kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 소과 2kg 10~13과(한입)")
            == "시나노골드(가정용) 소과 2kg(10-13과내외)")
    # 긴 등급 먼저: 중소과≠소과, 중대과랜덤≠대과, 중과≠중소과
    assert " 중소과 3kg" in kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 중소과 3kg")
    assert " 중대과랜덤 3kg" in kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 중대과 랜덤 3kg")
    assert " 중과 2.5kg" in kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 중과 2.5kg")
    assert " 대과 5kg" in kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "대과 5kg 1박스")
    # 제주다팜엔 '중대과'가 랜덤 하나뿐 → '중대과'·'랜덤'만 적혀도 중대과랜덤
    assert " 중대과랜덤 2kg" in kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 중대과 2kg")
    assert " 중대과랜덤 4kg" in kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "랜덤과 4kg")
    # 없는 중량은 변환 실패(발주서에선 needs_check로 표시)
    assert kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 소과 10kg") is None


def test_shinano_is_exclusive_from_other_products():
    assert kolrabi_order.convert_shinano_option(HWANGGEUM, "1박스 2kg (소)") is None
    hongro = "2026 햇 안동 고당도 홍로사과 아삭한 꿀사과 산지직송 제철 햇사과"
    assert kolrabi_order.convert_shinano_option(hongro, "1박스 소과 2kg") is None
    assert kolrabi_order.convert_hongro_option(COUPANG_PRODUCT, "1박스 소과 2kg 10~13과(한입)") is None
    assert kolrabi_order.convert_apple_option(COUPANG_PRODUCT, "1박스 소과 2kg 10~13과(한입)") is None
    assert not tomato_order.is_jbt_apple_order(COUPANG_PRODUCT, "1박스 소과 2kg 10~13과(한입)")
    assert order_routing.resolve_order_supplier(COUPANG_PRODUCT, "1박스 소과 2kg 10~13과(한입)") == "제주다팜"


def _delivery(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 3, r.get("order_no", f"OID{i}"))
        ws.cell(row, 11, r["product"])
        ws.cell(row, 12, r["option"])
        ws.cell(row, 23, r.get("qty", 1))
        ws.cell(row, 27, r.get("name", f"수취인{i}"))
        ws.cell(row, 28, r.get("phone", "010-0000-0000"))
        ws.cell(row, 29, "12345")
        ws.cell(row, 30, r.get("address", "서울시 어딘가"))
        ws.cell(row, 31, "")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_process_outputs_includes_shinano_file():
    data = _delivery([
        {"product": HWANGGEUM, "option": "1박스 2kg (소)", "name": "김광은", "order_no": "3103637528262"},
        {"product": COUPANG_PRODUCT, "option": "1박스 소과 2kg 10~13과(한입)", "name": "김광은", "order_no": "3103637528262"},
        {"product": COUPANG_PRODUCT, "option": "1박스 소과 10kg", "name": "미등록옵션"},
    ])
    results = kolrabi_order.process_outputs(data)
    shinano = [r for r in results if "시나노골드" in r[1]]
    # 황금향은 2026-10-16부터 별도 제주다팜 발주서(test_hwanggeum_jeju) — 시나노 파일엔 섞이지 않는다
    assert len(shinano) == 1 and {r[2]["product"] for r in results} == {"시나노골드사과(제주다팜)", "황금향(제주다팜)"}
    out, filename, stats = shinano[0]
    assert stats["product"] == "시나노골드사과(제주다팜)" and stats["total"] == 1
    assert any("미등록옵션" in n for n in stats["needs_check"])
    ws = load_workbook(BytesIO(out)).active
    assert (ws.cell(2, 2).value, ws.cell(2, 8).value, ws.cell(2, 9).value) == (
        "김광은", "시나노골드(가정용) 소과 2kg(10-13과내외)", "1")
    assert stats["options"][0]["orders"] == [{"order_id": "3103637528262", "quantity": 1}]


class _FakeTossClient:
    def __init__(self, orders):
        self._orders = orders

    async def get_orders(self, start_date=None, end_date=None, status=None):
        return self._orders


def test_toss_shinano_collected(monkeypatch):
    import app.toss.client as toss_module

    orders = [
        {"productName": COUPANG_PRODUCT, "optionName": "중과, 3kg, 1박스", "quantity": 1,
         "receiverName": "토스손님", "receiverPhone": "010-1111-2222", "orderProductStatus": "PAY_DONE"},
        {"productName": "[제주직송] 고당도 100% 제주 콜라비", "optionName": "정품 3kg", "quantity": 1,
         "receiverName": "콜라비손님", "receiverPhone": "010-1111-3333", "orderProductStatus": "PAY_DONE"},
    ]
    monkeypatch.setattr(toss_module, "toss_client", _FakeTossClient(orders), raising=False)
    result = asyncio.run(tomato_order.collect_toss_jejudapam_orders("2026-10-15", "2026-10-16"))
    assert [(e["name"], e["product"]) for e in result["shinano"]] == [
        ("토스손님", "시나노골드(가정용) 중과 3kg(9-12과내외)")]
    assert len(result["colrabi"]) == 1 and result["hongro"] == []


def test_tracking_semantic_keys_match_both_sides():
    ol = _semantic_option_keys("시나노골드(가정용) 소과 2kg(10-13과내외)", "")
    dl = _semantic_option_keys(COUPANG_PRODUCT, "1박스 소과 2kg 10~13과(한입)")
    assert "shinano:소과:2kg" in ol and "shinano:소과:2kg" in dl
    assert "shinano:중소과:2.5kg" in _semantic_option_keys("시나노골드(가정용) 중소과 2.5kg(11-12과내외)")
    assert "shinano:중대과랜덤:3kg" in _semantic_option_keys(COUPANG_PRODUCT, "1박스 중대과 3kg")


def _orderlist(rows: list[dict]) -> bytes:
    """제주다팜 회신: G 상품명·K 수령인·M 연락처·O 주소·Q 택배사·R 운송장."""
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 7, r["product"])
        ws.cell(row, 11, r["name"])
        ws.cell(row, 13, r.get("phone", "010-0000-0000"))
        ws.cell(row, 15, r.get("address", "서울시 어딘가"))
        ws.cell(row, 17, r.get("courier", "롯데택배"))
        ws.cell(row, 18, r["tracking"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracking_same_person_other_product_row_does_not_steal():
    """같은 사람(같은 주문번호)이 황금향+시나노골드를 주문 — 앞 행 황금향이 시나노 송장을 가로채면 안 된다."""
    delivery = _delivery([
        {"product": HWANGGEUM, "option": "1박스 2kg (소)", "name": "김광은", "order_no": "3103637528262"},
        {"product": COUPANG_PRODUCT, "option": "1박스 소과 2kg 10~13과(한입)", "name": "김광은", "order_no": "3103637528262"},
    ])
    orderlist = _orderlist([
        {"product": "시나노골드(가정용) 소과 2kg(10-13과내외)", "name": "김광은", "tracking": "254012345678"},
    ])
    out, _fn, stats = tracking_input.process(orderlist, delivery)
    ws = load_workbook(BytesIO(out)).active
    assert ws.cell(2, 5).value is None                                      # 황금향 행은 비움
    assert (ws.cell(3, 4).value, ws.cell(3, 5).value) == ("롯데택배", "254012345678")
    assert stats["filled"] == 1 and "다른 상품 행" in stats["skipped_names"]


def _real_format_orderlist(rows: list[tuple[str, str]]) -> bytes:
    """실제 제주다팜 회신처럼 E 상품명·G 수량·H 주문자·I 주문자연락처·N 우편번호까지 채운 회신."""
    wb = Workbook()
    ws = wb.active
    for i, (product, tracking) in enumerate(rows):
        r = i + 2
        ws.cell(r, 5, product)
        ws.cell(r, 7, 1)
        ws.cell(r, 8, "식품애착")
        ws.cell(r, 9, "010-5700-7756")
        ws.cell(r, 11, "메리")
        ws.cell(r, 13, "010-0000-0000")
        ws.cell(r, 14, "12345")
        ws.cell(r, 15, "서울시 어딘가")
        ws.cell(r, 17, "롯데택배")
        ws.cell(r, 18, tracking)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracking_potato_vendor_grade_name_still_fills_with_two_rows():
    """리뷰 회귀(2026-10-16): 제주다팜은 홍감자를 '홍감자 대_3KG'(=우리 중 3kg)로 회신해 등급 키가 안 맞는다.
    같은 사람이 DeliveryList에 두 줄 이상이어도 같은 상품군이면 '다른 상품 행'으로 건너뛰면 안 된다."""
    potato = "[여름한정] 국내산 햇 홍감자 카스테라 햇감자"
    # 홍감자 2건(같은 사람)
    delivery = _delivery([
        {"product": potato, "option": "1박스 3kg(중)", "name": "메리", "order_no": "P1"},
        {"product": potato, "option": "1박스 3kg(중)", "name": "메리", "order_no": "P2"},
    ])
    reply = _real_format_orderlist([("홍감자 대_3KG (개당 80-110g)", "T1"), ("홍감자 대_3KG (개당 80-110g)", "T2")])
    out, _fn, stats = tracking_input.process(reply, delivery)
    ws = load_workbook(BytesIO(out)).active
    assert [ws.cell(r, 5).value for r in (2, 3)] == ["T1", "T2"] and stats["filled"] == 2
    # 홍감자 + 다른 거래처 상품(고구마) — 홍감자 행은 채우고 고구마 행은 비움
    delivery = _delivery([
        {"product": "해남 꿀고구마", "option": "1박스 3kg", "name": "메리", "order_no": "G1"},
        {"product": potato, "option": "1박스 3kg(중)", "name": "메리", "order_no": "P1"},
    ])
    reply = _real_format_orderlist([("홍감자 대_3KG (개당 80-110g)", "T1")])
    out, _fn, stats = tracking_input.process(reply, delivery)
    ws = load_workbook(BytesIO(out)).active
    assert [ws.cell(r, 5).value for r in (2, 3)] == [None, "T1"]
    # 홍감자 + 콜라비(회신 순서 반대) — 각자 맞는 송장
    delivery = _delivery([
        {"product": potato, "option": "1박스 3kg(중)", "name": "메리", "order_no": "P1"},
        {"product": "[제주직송] 고당도 100% 제주 콜라비", "option": "정품 3kg", "name": "메리", "order_no": "K1"},
    ])
    reply = _real_format_orderlist([("콜라비 정품 3kg", "K1"), ("홍감자 대_3KG (개당 80-110g)", "P1")])
    out, _fn, _stats = tracking_input.process(reply, delivery)
    ws = load_workbook(BytesIO(out)).active
    assert [ws.cell(r, 5).value for r in (2, 3)] == ["P1", "K1"]


def test_tracking_same_person_same_product_without_reply_options_still_fills():
    """회신에 옵션 정보가 없으면 종전처럼 전화로 매칭(2026-07-14 동일인 다건 사고 보호)."""
    delivery = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1박스 소과 2kg", "name": "동일인", "order_no": "A1"},
        {"product": COUPANG_PRODUCT, "option": "1박스 소과 2kg", "name": "동일인", "order_no": "A2"},
    ])
    wb = Workbook()
    ws = wb.active
    for i, t in enumerate(("111", "222")):
        ws.cell(i + 2, 11, "동일인")
        ws.cell(i + 2, 13, "010-0000-0000")
        ws.cell(i + 2, 15, "서울시 어딘가")
        ws.cell(i + 2, 18, t)
    buf = BytesIO()
    wb.save(buf)
    out, _fn, stats = tracking_input.process(buf.getvalue(), delivery)
    assert stats["filled"] == 2
