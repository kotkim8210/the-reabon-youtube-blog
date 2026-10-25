"""감홍사과 제주다팜 발주·토스 수집·운송장 (2026-10-26 신규).

판매옵션 12종(adminplus pcode=10001367 실측): '감홍사과(가정용 A등급) {대과|중과} {1.5|2|2.5|3|4|5}kg(과수)'.
쿠팡 '[산지직송] 고당도 경북 문경 감홍사과 가정용 A등급' 옵션 '1박스 1.5kg (대과)'(DeliveryList 2026-10-26 실측, 83건).
★ '감홍사과'에 '홍사과'가 들어 있어 종전엔 홍로사과(가을햇사과)로 둔갑해 발주될 뻔했다 — 배타 판정 회귀 포함.
"""

import asyncio
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import order_routing
from app.processors import kolrabi_order, tomato_order, tracking_input
from app.processors.tracking_input import _semantic_option_keys

COUPANG_PRODUCT = "[산지직송] 고당도 경북 문경 감홍사과 가정용 A등급"
HONGRO = "2026 햇 안동 고당도 홍로사과 아삭한 꿀사과 산지직송 제철 햇사과"


def test_all_12_options_map():
    for grade in ("대과", "중과"):
        for kg in ("1.5", "2", "2.5", "3", "4", "5"):
            got = kolrabi_order.convert_gamhong_option(COUPANG_PRODUCT, f"1박스 {kg}kg ({grade})")
            assert got and got.startswith(f"감홍사과(가정용 A등급) {grade} {kg}kg("), (grade, kg, got)
    assert len(set(kolrabi_order._GAMHONG_JEJU_OPTIONS.values())) == 12


def test_real_options():
    cases = {
        "1박스 1.5kg (대과)": "감홍사과(가정용 A등급) 대과 1.5kg(2-4과내)",
        "1박스 1.5kg (중과)": "감홍사과(가정용 A등급) 중과 1.5kg(4-5과내)",
        "1박스 3kg (중과)": "감홍사과(가정용 A등급) 중과 3kg(9-13과내외)",
        "1박스 5kg (대과)": "감홍사과(가정용 A등급) 대과 5kg(8-14과내외)",
        "1박스 5kg (중과)": "감홍사과(가정용 A등급) 중과 5kg(15-20과내외)",
    }
    for option, expected in cases.items():
        assert kolrabi_order.convert_gamhong_option(COUPANG_PRODUCT, option) == expected, option
    # 'A등급'은 등급 토큰이 아니다, 소과는 공급처에 없다 → 변환 실패(needs_check)
    assert kolrabi_order.convert_gamhong_option(COUPANG_PRODUCT, "1박스 2kg") is None
    assert kolrabi_order.convert_gamhong_option(COUPANG_PRODUCT, "1박스 2kg (소과)") is None


def test_gamhong_never_becomes_hongro_and_vice_versa():
    """'감홍사과' ⊃ '홍사과' — 홍로 판정·변환·토스·운송장 키 어디서도 섞이면 안 된다."""
    assert kolrabi_order.is_jeju_hongro_order(COUPANG_PRODUCT, "1박스 1.5kg (대과)") is False
    assert kolrabi_order.convert_hongro_option(COUPANG_PRODUCT, "1박스 1.5kg (대과)") is None
    assert kolrabi_order.convert_gamhong_option(HONGRO, "1박스 대과 5kg(16-18과)") is None
    assert kolrabi_order.convert_hongro_option(HONGRO, "1박스 대과 5kg(16-18과)") == "가을햇사과(홍사과) 가정용 대과 포장재포함 5kg(16-18과)"
    assert order_routing.resolve_order_supplier(COUPANG_PRODUCT, "1박스 1.5kg (대과)") == "제주다팜"
    keys = _semantic_option_keys(COUPANG_PRODUCT, "1박스 1.5kg (대과)")
    assert "gamhong:대과:1.5kg" in keys and not any(k.startswith("hongro:") for k in keys)
    assert not any(k.startswith("gamhong:") for k in _semantic_option_keys(HONGRO, "1박스 대과 5kg"))


def _delivery(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 3, r.get("order_no", f"OID{i}"))
        ws.cell(row, 11, r["product"])
        ws.cell(row, 12, r["option"])
        ws.cell(row, 23, r.get("qty", 1))
        ws.cell(row, 27, r.get("name", f"수취인{i}"))
        ws.cell(row, 28, r.get("phone", "010-0000-0000"))
        ws.cell(row, 29, "12345")
        ws.cell(row, 30, r.get("address", "서울시 어딘가"))
        ws.cell(row, 31, "")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_process_outputs_separates_gamhong_from_hongro():
    data = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1박스 3kg (중과)", "name": "김남구"},
        {"product": COUPANG_PRODUCT, "option": "1박스 5kg (대과)", "name": "양지안"},
        {"product": HONGRO, "option": "1박스 대과 5kg(16-18과)", "name": "유성자"},
    ])
    results = {r[2]["product"]: r for r in kolrabi_order.process_outputs(data)}
    assert set(results) == {"감홍사과(제주다팜)", "홍로사과(제주다팜)"}
    out, filename, stats = results["감홍사과(제주다팜)"]
    assert "감홍사과" in filename and stats["total"] == 2 and "needs_check" not in stats
    ws = load_workbook(BytesIO(out)).active
    assert [(ws.cell(r, 2).value, ws.cell(r, 8).value) for r in (2, 3)] == [
        ("김남구", "감홍사과(가정용 A등급) 중과 3kg(9-13과내외)"), ("양지안", "감홍사과(가정용 A등급) 대과 5kg(8-14과내외)")]
    _o, _f, hongro_stats = results["홍로사과(제주다팜)"]
    assert hongro_stats["total"] == 1   # 감홍 2건이 홍로 파일에 섞이지 않는다


class _FakeTossClient:
    def __init__(self, orders):
        self._orders = orders

    async def get_orders(self, start_date=None, end_date=None, status=None):
        return self._orders


def test_toss_gamhong_collected_not_as_hongro(monkeypatch):
    import app.toss.client as toss_module

    orders = [{"productName": COUPANG_PRODUCT, "optionName": "대과, 1.5kg, 1박스", "quantity": 1,
               "receiverName": "토스손님", "receiverPhone": "010-1111-2222", "orderProductStatus": "PAY_DONE"}]
    monkeypatch.setattr(toss_module, "toss_client", _FakeTossClient(orders), raising=False)
    result = asyncio.run(tomato_order.collect_toss_jejudapam_orders("2026-10-25", "2026-10-26"))
    assert [(e["name"], e["product"]) for e in result["gamhong"]] == [("토스손님", "감홍사과(가정용 A등급) 대과 1.5kg(2-4과내)")]
    assert result["hongro"] == []


def test_tracking_fill_same_person_two_sizes():
    delivery = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1박스 1.5kg (중과)", "name": "한선영"},
        {"product": COUPANG_PRODUCT, "option": "1박스 5kg (대과)", "name": "한선영"},
    ])
    wb = Workbook()
    ws = wb.active
    for i, (product, tracking) in enumerate((("감홍사과(가정용 A등급) 대과 5kg(8-14과내외)", "T5"),
                                             ("감홍사과(가정용 A등급) 중과 1.5kg(4-5과내)", "T15"))):
        r = i + 2
        ws.cell(r, 5, product)
        ws.cell(r, 11, "한선영")
        ws.cell(r, 13, "010-0000-0000")
        ws.cell(r, 15, "서울시 어딘가")
        ws.cell(r, 17, "롯데택배")
        ws.cell(r, 18, tracking)
    buf = BytesIO()
    wb.save(buf)
    out, _fn, stats = tracking_input.process(buf.getvalue(), delivery)
    ws = load_workbook(BytesIO(out)).active
    assert [ws.cell(r, 5).value for r in (2, 3)] == ["T15", "T5"] and stats["filled"] == 2
