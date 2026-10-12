"""미출고 자동 재발주 — 영업일 2일↑ 지난 발주분이 아직 DeliveryList/토스에 남아 있으면 다시 넣는다 (2026-10-12).

사고: 10/8(목, 해달 휴무) 09:21에 만든 호박고구마 발주서 2건(김영옥·정복순)이 처리되지 않았는데,
10/12(월) 발주서에서 '이전 발주분'으로 자동 제외돼 누락. 경고(needs_check)만 띄우던 설계를
자동 재발주 + '미출고 재발주' 표시로 바꿨다.
"""

import asyncio
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.processors import goguma_auto, issued_orders

KEY_A = "15103446848104|1박스2kg(한입:에프용/아이간식)"   # 10/8 발주분 → 재발주 대상
KEY_B = "20103446848105|1박스5kg(중상)"                  # 어제 발주분 → 계속 제외


def _delivery(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["헤더"] * 31)
    for order_no, option, name in rows:
        row = [""] * 31
        row[2], row[10], row[11], row[26] = order_no, "국내산 해남 황토밭 햇 호풍미 호박고구마 판매1위", option, name
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_delivery_filter_keeps_stale_rows_and_reports_them():
    data = _delivery([
        ("15103446848104", "1박스 2kg(한입:에프용/아이간식)", "김영옥"),   # 재발주
        ("20103446848105", "1박스 5kg (중상)", "어제손님"),               # 제외
        ("30103446848106", "1박스 3kg (중상)", "신규손님"),               # 그대로
    ])
    dates = {KEY_A: "2026-10-08", KEY_B: "2026-10-09"}
    stale = issued_orders.stale_issued_keys(dates, "2026-10-12")
    exclude = set(dates) - stale
    dup_names, dup_keys, reissued = [], [], []
    out, skipped = issued_orders.filter_delivery_by_issued(
        data, exclude, skipped_names=dup_names, skipped_keys=dup_keys, reissue_keys=stale, reissued=reissued
    )
    assert skipped == 1 and dup_names == ["어제손님"] and dup_keys == [KEY_B]
    assert reissued == [("김영옥", KEY_A)]
    ws = load_workbook(BytesIO(out)).active
    assert [ws.cell(r, 27).value for r in range(2, ws.max_row + 1)] == ["김영옥", "신규손님"]


def test_entries_filter_reissues_stale_toss_orders():
    entries = [
        {"order_id": "288697616", "product": "호박고구마 2Kg (한입)", "name": "토스옛손님"},
        {"order_id": "288697617", "product": "호박고구마 2Kg (한입)", "name": "토스어제손님"},
        {"order_id": "288697618", "product": "호박고구마 5Kg (중상)", "name": "토스신규"},
    ]
    dates = {"288697616|호박고구마2Kg(한입)": "2026-10-06", "288697617|호박고구마2Kg(한입)": "2026-10-09"}
    stale = issued_orders.stale_issued_keys(dates, "2026-10-12")
    names, reissued = [], []
    kept, dup = issued_orders.filter_entries_by_issued(entries, set(dates) - stale, skipped_names=names,
                                                       reissue_keys=stale, reissued=reissued)
    assert dup == 1 and names == ["토스어제손님"]
    assert [e["name"] for e in kept] == ["토스옛손님", "토스신규"]
    assert reissued == [("토스옛손님", "288697616|호박고구마2Kg(한입)")]


def test_goguma_api_path_matches_composite_keys(monkeypatch):
    """고구마 페이지 API 수집: 기록은 '주문번호|품목'인데 주문번호만 비교해 제외가 안 걸리던 구멍 회귀."""
    async def fake_collect(_from, _to):
        return [
            {"name": "어제발주", "phone": "", "zipcode": "", "address": "서울", "quantity": 1,
             "product": "꿀고구마 3Kg (중상)", "memo": "", "order_id": "8103454146245"},
            {"name": "신규", "phone": "", "zipcode": "", "address": "서울", "quantity": 1,
             "product": "꿀고구마 5Kg (한입)", "memo": "", "order_id": "9103454146246"},
        ]
    monkeypatch.setattr(goguma_auto, "collect_orders", fake_collect)
    names, reissued = [], []
    _b, _fn, stats = asyncio.run(goguma_auto.process_from_api(
        "2026-10-11", "2026-10-12", exclude_order_ids={"8103454146245|꿀고구마3Kg(중상)"},
        skipped_names=names, reissue_keys=set(), reissued=reissued,
    ))
    assert stats["total"] == 1 and stats["duplicate_skipped"] == 1
    assert names == ["어제발주"] and reissued == []


# ── 리뷰 반영 (2026-10-13) ──
def test_effective_issue_date_shifts_late_and_weekend_issues():
    eff = issued_orders.effective_issue_date
    assert eff("2026-10-08", "2026-10-08 00:21:19") == "2026-10-08"   # 목 09:21 KST 생성 → 당일
    assert eff("2026-10-12", "2026-10-12 07:17:00") == "2026-10-13"   # 월 16:17 KST(마감 후) → 화요일 접수
    assert eff("2026-10-11", "2026-10-11 02:00:00") == "2026-10-12"   # 일요일 생성 → 월요일 접수
    assert eff("2026-10-09", "2026-10-09 01:29:00") == "2026-10-09"   # 금 10:29 → 당일
    assert eff("2026-10-09", "2026-10-09 01:30:00") == "2026-10-12"   # 금 10:30 → 다음 영업일 월
    assert eff("2026-10-08", None) == "2026-10-08"
    assert eff("깨진값", "x") == "깨진값"
    # 월 16:17 발주 → 수요일엔 아직 1영업일(제외 유지), 목요일에 재발주
    assert issued_orders.business_days_since(eff("2026-10-12", "2026-10-12 07:17:00"), "2026-10-14") == 1
    assert issued_orders.business_days_since(eff("2026-10-12", "2026-10-12 07:17:00"), "2026-10-15") == 2


def test_entries_filter_matches_old_transformed_key_and_new_raw_key():
    """쿠팡 API 경로: 과거 기록(변환명 키)과 새 기록(원문 키) 둘 다 제외에 걸려야 한다."""
    entry = {"order_id": "8103454146245", "option": "1박스 황금 꿀고구마 3Kg (중상)",
             "product": "꿀고구마 3Kg (중상)", "name": "손님"}
    for key in ("8103454146245|꿀고구마3Kg(중상)", "8103454146245|1박스황금꿀고구마3Kg(중상)"):
        kept, dup = issued_orders.filter_entries_by_issued([dict(entry)], {key})
        assert (kept, dup) == ([], 1), key


def test_alwayz_rows_with_tracking_are_skipped():
    from app.processors import goguma_order
    wb = Workbook()
    ws = wb.active
    ws.append(["헤더"] * 24)
    shipped = [""] * 24; shipped[18] = "출고된손님"; shipped[5] = "0. 해남 황금 꿀고구마: 3Kg 중상"; shipped[22] = "463300000001"
    pending = [""] * 24; pending[18] = "대기손님"; pending[5] = "0. 해남 황금 꿀고구마: 2Kg 한입"
    ws.append(shipped); ws.append(pending)
    buf = BytesIO(); wb.save(buf)
    names = [e["name"] for e in goguma_order.parse_alwayz_entries(buf.getvalue())]
    assert names == ["대기손님"]


def test_widen_start_reaches_back_to_stale_issue_date(monkeypatch):
    from app import main as app_main
    monkeypatch.setattr(app_main, "_today_kst", lambda: "2026-10-12")
    dates = {"A|x": "2026-10-08", "B|x": "2026-09-01"}
    assert app_main._widen_start("2026-10-11", dates, {"A|x"}) == "2026-10-08"
    assert app_main._widen_start("2026-10-11", dates, {"B|x"}) == "2026-09-28"   # 14일 상한
    assert app_main._widen_start("2026-10-11", dates, set()) == "2026-10-11"
    assert app_main._widen_start("2026-10-01", dates, {"A|x"}) == "2026-10-01"    # 이미 더 넓으면 그대로
