"""태추단감 제주다팜 발주·토스 수집·운송장 (2026-10-19 신규).

판매옵션 12종: '태추단감 {대과|중과|소과} {1|1.5|2|3}kg (과수)'.
쿠팡 '경남 태추단감 아삭달콤 제철 배단감' 옵션 '1박스 1kg (중과)'(DeliveryList 2026-10-19 실측, 46건).
"""

import asyncio
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import order_routing
from app.processors import kolrabi_order, tomato_order, tracking_input
from app.processors.tracking_input import _semantic_option_keys

COUPANG_PRODUCT = "경남 태추단감 아삭달콤 제철 배단감"


def test_all_12_options_map():
    for grade in ("대과", "중과", "소과"):
        for kg in ("1", "1.5", "2", "3"):
            got = kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, f"1박스 {kg}kg ({grade})")
            assert got and got.startswith(f"태추단감 {grade} {kg}kg ("), (grade, kg, got)
    assert len(set(kolrabi_order._TAECHU_JEJU_OPTIONS.values())) == 12


def test_real_options_and_edge_cases():
    assert kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, "1박스 1kg (중과)") == "태추단감 중과 1kg (4-5과)"
    assert kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, "1박스 1.5kg (중과)") == "태추단감 중과 1.5kg (6-7과)"
    assert kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, "1박스 3kg (중과)") == "태추단감 중과 3kg (13-15과)"
    # 한 글자 등급
    assert kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, "1박스 2kg (대)") == "태추단감 대과 2kg (7-8과)"
    # 등급 없음·없는 중량은 변환 실패(발주서 needs_check)
    assert kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, "1박스 2kg") is None
    assert kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, "1박스 5kg (중과)") is None
    assert kolrabi_order.convert_taechu_option(COUPANG_PRODUCT, "대용량 2kg") is None
    # 다른 단감·다른 제주다팜 품목과 섞이지 않음
    assert kolrabi_order.convert_taechu_option("부유단감 가정용", "1박스 2kg (중과)") is None
    assert kolrabi_order.convert_hwanggeum_option(COUPANG_PRODUCT, "1박스 1kg (중과)") is None
    assert kolrabi_order.convert_shinano_option(COUPANG_PRODUCT, "1박스 1kg (중과)") is None
    assert order_routing.resolve_order_supplier(COUPANG_PRODUCT, "1박스 1kg (중과)") == "제주다팜"


def _delivery(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 3, r.get("order_no", f"OID{i}"))
        ws.cell(row, 11, r["product"])
        ws.cell(row, 12, r["option"])
        ws.cell(row, 23, r.get("qty", 1))
        ws.cell(row, 27, r.get("name", f"수취인{i}"))
        ws.cell(row, 28, r.get("phone", "010-0000-0000"))
        ws.cell(row, 29, "12345")
        ws.cell(row, 30, r.get("address", "서울시 어딘가"))
        ws.cell(row, 31, "")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_process_outputs_includes_taechu_file():
    data = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1박스 1kg (중과)", "name": "이춘우", "order_no": "4103690732205"},
        {"product": COUPANG_PRODUCT, "option": "1박스 3kg (중과)", "name": "김용미", "order_no": "20103690831278"},
        {"product": "2026 햇 안동 고당도 홍로사과 아삭한 꿀사과 산지직송 제철 햇사과", "option": "1박스 대과 5kg(16-18과)"},
    ])
    results = {r[2]["product"]: r for r in kolrabi_order.process_outputs(data)}
    assert set(results) == {"태추단감(제주다팜)", "홍로사과(제주다팜)"}
    out, filename, stats = results["태추단감(제주다팜)"]
    assert "태추단감" in filename and stats["total"] == 2 and "needs_check" not in stats
    ws = load_workbook(BytesIO(out)).active
    assert [(ws.cell(r, 2).value, ws.cell(r, 8).value) for r in (2, 3)] == [
        ("이춘우", "태추단감 중과 1kg (4-5과)"), ("김용미", "태추단감 중과 3kg (13-15과)")]


class _FakeTossClient:
    def __init__(self, orders):
        self._orders = orders

    async def get_orders(self, start_date=None, end_date=None, status=None):
        return self._orders


def test_toss_taechu_collected(monkeypatch):
    import app.toss.client as toss_module

    orders = [{"productName": COUPANG_PRODUCT, "optionName": "중과, 2kg, 1박스", "quantity": 1,
               "receiverName": "토스손님", "receiverPhone": "010-1111-2222", "orderProductStatus": "PAY_DONE"}]
    monkeypatch.setattr(toss_module, "toss_client", _FakeTossClient(orders), raising=False)
    result = asyncio.run(tomato_order.collect_toss_jejudapam_orders("2026-10-18", "2026-10-19"))
    assert [(e["name"], e["product"]) for e in result["taechu"]] == [("토스손님", "태추단감 중과 2kg (9-10과)")]


def test_semantic_keys_and_tracking_fill():
    assert "taechu:중과:1kg" in _semantic_option_keys("태추단감 중과 1kg (4-5과)", "")
    assert "taechu:중과:1kg" in _semantic_option_keys(COUPANG_PRODUCT, "1박스 1kg (중과)")
    assert "taechu:중과:1.5kg" in _semantic_option_keys(COUPANG_PRODUCT, "1박스 1.5kg (중과)")

    delivery = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1박스 1kg (중과)", "name": "이춘우"},
        {"product": COUPANG_PRODUCT, "option": "1박스 3kg (중과)", "name": "이춘우"},   # 같은 사람 2건, 회신 순서 반대
    ])
    wb = Workbook()
    ws = wb.active
    for i, (product, tracking) in enumerate((("태추단감 중과 3kg (13-15과)", "T3"), ("태추단감 중과 1kg (4-5과)", "T1"))):
        r = i + 2
        ws.cell(r, 5, product)
        ws.cell(r, 11, "이춘우")
        ws.cell(r, 13, "010-0000-0000")
        ws.cell(r, 15, "서울시 어딘가")
        ws.cell(r, 17, "롯데택배")
        ws.cell(r, 18, tracking)
    buf = BytesIO()
    wb.save(buf)
    out, _fn, stats = tracking_input.process(buf.getvalue(), delivery)
    ws = load_workbook(BytesIO(out)).active
    assert [(ws.cell(r, 4).value, ws.cell(r, 5).value) for r in (2, 3)] == [("롯데택배", "T1"), ("롯데택배", "T3")]
    assert stats["filled"] == 2
