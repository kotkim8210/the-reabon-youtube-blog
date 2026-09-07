"""비셀러 LA한입갈비 운송장번호 입력 (2026-09-07 신규)."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.processors import biseller_tracking
from app.processors.tracking_match import coupang_courier_name

_REPLY_HEADERS = [
    "순번", "발주일", "수취인명", "수취인연락처", "우편번호", "주소",
    "상품명 (비셀러 상품명)", "수량", "배송메세지", "주문자명", "주문자연락처",
    "택배사", "송장번호",
]


def _reply(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(_REPLY_HEADERS)
    for i, r in enumerate(rows):
        ws.append([
            i + 1, "2026-09-06", r["name"], r.get("phone", "0502-1111-2222"), "12345",
            r.get("address", "서울시 어딘가 1-2"), r.get("product", "양념LA한입갈비 800g+800g (800G*2세트)"),
            1, "문 앞", "(주)아이티소프트", "010-5700-7756",
            r.get("courier", "롯데(현대)택배"), r.get("tracking", ""),
        ])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _delivery(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 3, "주문번호")
    ws.cell(1, 5, "운송장번호")
    ws.cell(1, 27, "수취인이름")
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 3, r.get("order_no", "OID" + str(i)))
        ws.cell(row, 11, "한입 LA갈비 양념갈비")
        ws.cell(row, 12, r.get("option", "800g 2개"))
        ws.cell(row, 27, r["name"])
        ws.cell(row, 28, r.get("phone", "0502-1111-2222"))
        ws.cell(row, 30, r.get("address", "서울시 어딘가 1-2"))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_lotte_variant_is_normalized_for_coupang():
    """비셀러는 '롯데(현대)택배'로 회신 — 쿠팡은 '롯데택배'만 인식한다."""
    assert coupang_courier_name("롯데(현대)택배") == "롯데택배"
    assert coupang_courier_name("롯데택배") == "롯데택배"
    assert coupang_courier_name("한진택배") == "한진택배"


def test_fills_tracking_from_two_reply_files():
    delivery = _delivery([{"name": "배수영"}, {"name": "조희래"}, {"name": "제원희"}])
    reply1 = _reply([
        {"name": "배수영", "tracking": "411722399325"},
        {"name": "조희래", "tracking": "411722399561"},
    ])
    reply2 = _reply([{"name": "제원희", "tracking": "411722410665"}])
    out, filename, stats = biseller_tracking.process(delivery, [reply1, reply2])
    assert stats["filled"] == 3 and stats["skipped"] == 0
    assert filename.startswith("DeliveryList_비셀러_운송장입력완료_")
    ws = load_workbook(BytesIO(out)).active
    assert [ws.cell(r, 5).value for r in (2, 3, 4)] == ["411722399325", "411722399561", "411722410665"]
    assert ws.cell(2, 4).value == "롯데택배"


def test_unmatched_reply_row_is_reported():
    """이벤트 경품처럼 쿠팡 주문이 없는 회신 행은 조용히 버리지 않는다."""
    delivery = _delivery([{"name": "배수영"}])
    reply = _reply([
        {"name": "배수영", "tracking": "411722399325"},
        {"name": "이벤트당첨자", "tracking": "411722410665"},
    ])
    _out, _fn, stats = biseller_tracking.process(delivery, [reply])
    assert stats["filled"] == 1 and stats["skipped"] == 1
    assert "이벤트당첨자" in stats["needs_check"][0]


def test_same_name_is_split_by_safe_phone_number():
    delivery = _delivery([
        {"name": "김철수", "phone": "0502-1111-1111"},
        {"name": "김철수", "phone": "0502-2222-2222"},
    ])
    reply = _reply([
        {"name": "김철수", "phone": "0502-2222-2222", "tracking": "411722400001"},
        {"name": "김철수", "phone": "0502-1111-1111", "tracking": "411722400002"},
    ])
    out, _fn, stats = biseller_tracking.process(delivery, [reply])
    assert stats["filled"] == 2
    ws = load_workbook(BytesIO(out)).active
    assert ws.cell(2, 5).value == "411722400002"
    assert ws.cell(3, 5).value == "411722400001"


def test_wrong_file_in_delivery_slot_raises_korean_error():
    reply = _reply([{"name": "배수영", "tracking": "411722399325"}])
    try:
        biseller_tracking.process(reply, [reply])
    except ValueError as exc:
        assert "DeliveryList" in str(exc)
    else:
        raise AssertionError("잘못된 파일이 통과되면 안 된다")


def test_reply_without_tracking_raises():
    delivery = _delivery([{"name": "배수영"}])
    reply = _reply([{"name": "배수영", "tracking": ""}])
    try:
        biseller_tracking.process(delivery, [reply])
    except ValueError as exc:
        assert "송장번호" in str(exc)
    else:
        raise AssertionError("송장 없는 회신이 통과되면 안 된다")
