"""라이브 이벤트 당첨자 → 제주다팜: 감홍사과 등 등급별 판매옵션 품목 (2026-10-26).

당첨자 CSV 경품명은 '감홍사과1.5kg'처럼 등급 없이 온다. 종전엔 경품명 원문이 발주서에 그대로 나가
거래처가 대과/중과를 알 수 없었다 → 등급 없으면 백도 경품과 같은 기본값 '중과'로 판매옵션명을 붙인다.
"""

from io import BytesIO

from openpyxl import load_workbook

from app.processors import event_order

HEADER = "경품명,별명,개인 식별 정보 확인,이름,연락처,주소,주문 아이디,구매 금액,환불/취소 금액,환불/날짜/시간\n"


def _csv(rows: list[str]) -> bytes:
    return (HEADER + "\n".join(rows) + "\n").encode("utf-8-sig")


def test_gradeless_gamhong_prize_becomes_junggwa_option():
    assert event_order.event_product_name("감홍사과1.5kg") == "감홍사과(가정용 A등급) 중과 1.5kg(4-5과내)"
    assert event_order.event_product_name("감홍사과 3kg") == "감홍사과(가정용 A등급) 중과 3kg(9-13과내외)"
    # 경품에 등급이 적혀 있으면 그 등급
    assert event_order.event_product_name("감홍사과 대과 1.5kg") == "감홍사과(가정용 A등급) 대과 1.5kg(2-4과내)"
    # 공급처에 없는 중량은 원문 유지(엉뚱한 옵션으로 둔갑 금지)
    assert event_order.event_product_name("감홍사과 1kg") == "감홍사과 1kg"
    assert event_order.event_supplier("감홍사과1.5kg") == "jejudapam"


def test_other_graded_jeju_products_and_existing_prizes_unchanged():
    assert event_order.event_product_name("태추단감 1kg") == "태추단감 중과 1kg (4-5과)"
    assert event_order.event_product_name("황금향 2kg") == "하우스 정품 황금향 중과(150~220g) 2kg"
    assert event_order.event_product_name("시나노골드 2kg") == "시나노골드(가정용) 중과 2kg(6-7과내외)"
    # 기존 경품 규칙 그대로
    assert event_order.event_product_name("백도 딱딱이 복숭아 1kg") == "딱딱이 복숭아 중과 2kg"
    assert event_order.event_product_name("홍로사과 대과 5kg") == "가을햇사과(홍사과) 가정용 대과 포장재포함 5kg(16-18과)"
    assert event_order.event_product_name("성주참외 2kg") == "성주참외 2kg"


def test_process_real_shaped_csv():
    data = _csv([
        "감홍사과1.5kg,별명1,2026/10/25 22:01,테스트가,010-0000-0001,대구광역시 북구 테스트로 1,21103842199632,34400,0,",
        "감홍사과1.5kg,별명2,2026/10/25 22:01,테스트나,010-0000-0002,전라남도 무안군 테스트로 2,23103842569825,34900,0,",
    ])
    results = event_order.process(data)
    assert len(results) == 1
    out, filename, stats = results[0]
    assert filename.startswith("이벤트당첨_제주다팜_감홍사과_발주(") and stats["winners"] == 2
    ws = load_workbook(BytesIO(out)).active
    assert [(ws.cell(r, 2).value, ws.cell(r, 8).value) for r in (2, 3)] == [
        ("테스트가", "감홍사과(가정용 A등급) 중과 1.5kg(4-5과내)"),
        ("테스트나", "감홍사과(가정용 A등급) 중과 1.5kg(4-5과내)"),
    ]
