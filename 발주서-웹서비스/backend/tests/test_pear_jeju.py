"""프리미엄 햇 배 선물세트 제주다팜 발주·운송장 (2026-09-20 신규).

쿠팡 실측(2026-09-20 DeliveryList): K='프리미엄 햇 배 선물세트 보자기 포함 큼직한 특대과 5kg', L='1세트'
→ 제주다팜 구성상품명 '햇 배 선물세트 특대과 5kg (6~8과) + 보자기 동봉'.
발송은 CJ대한통운 — 쿠팡 D열은 'CJ 대한통운'(띄어쓰기)만 인식한다.
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.processors import kolrabi_order
from app.processors.tracking_input import _semantic_option_keys, process as tracking_process
from app.processors.tracking_match import COUPANG_CJ_NAME, coupang_courier_name, normalize_courier_name

COUPANG_PRODUCT = "프리미엄 햇 배 선물세트 보자기 포함 큼직한 특대과 5kg"
JEJU_NAME = "햇 배 선물세트 특대과 5kg (6~8과) + 보자기 동봉"


# ── 발주명 변환 ──
def test_real_coupang_row_maps_to_jeju_name():
    assert kolrabi_order.is_jeju_pear_order(COUPANG_PRODUCT, "1세트")
    assert kolrabi_order.convert_pear_option(COUPANG_PRODUCT, "1세트") == JEJU_NAME


def test_all_8_vendor_options_map():
    """대과/특대과 × 5·7.5kg × 보자기 유무 — adminplus 구성상품 8종 실측명 그대로."""
    expected = {
        ("보자기 포함", "특대과", "5"): "햇 배 선물세트 특대과 5kg (6~8과) + 보자기 동봉",
        ("보자기 포함", "특대과", "7.5"): "햇 배 선물세트 특대과 7.5kg (9~11과) + 보자기 동봉",
        ("보자기 포함", "대과", "5"): "햇 배 선물세트 대과 5kg (9~12과) + 보자기 동봉",
        ("보자기 포함", "대과", "7.5"): "햇 배 선물세트 대과 7.5kg (13~16과) + 보자기 동봉",
        ("", "특대과", "5"): "햇 배 선물세트 특대과 5kg (5~8과)",
        ("", "특대과", "7.5"): "햇 배 선물세트 특대과 7.5kg (9-12과)",
        ("", "대과", "5"): "햇 배 선물세트 대과 5kg (9~12과)",
        ("", "대과", "7.5"): "햇 배 선물세트 대과 7.5kg (13~16과)",
    }
    for (bojagi, grade, kg), jeju in expected.items():
        product = f"프리미엄 햇 배 선물세트 {bojagi} 큼직한 {grade} {kg}kg".replace("  ", " ")
        assert kolrabi_order.convert_pear_option(product, "1세트") == jeju, (bojagi, grade, kg)


def test_pear_is_exclusive_with_other_products():
    # 배 선물세트는 홍로·청사과 변환기에 안 잡히고, 사과 선물세트는 배 변환기에 안 잡힌다
    assert kolrabi_order.convert_hongro_option(COUPANG_PRODUCT, "1세트") is None
    assert kolrabi_order.convert_apple_option(COUPANG_PRODUCT, "1세트") is None
    assert kolrabi_order.convert_pear_option("프리미엄 햇 사과 선물세트 보자기 포함 특대과 5kg", "1세트") is None
    assert kolrabi_order.convert_pear_option("2026 햇 안동 고당도 홍로사과", "1박스 대과 5kg") is None
    # 취급 안 하는 중량은 None (matcher는 True → needs_check로 올라간다)
    assert kolrabi_order.is_jeju_pear_order(COUPANG_PRODUCT.replace("5kg", "10kg"), "1세트")
    assert kolrabi_order.convert_pear_option(COUPANG_PRODUCT.replace("5kg", "10kg"), "1세트") is None


def _delivery(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 3, r.get("order_no", f"OID{i}"))
        ws.cell(row, 11, r["product"])
        ws.cell(row, 12, r["option"])
        ws.cell(row, 23, r.get("qty", 1))
        ws.cell(row, 27, r.get("name", f"수취인{i}"))
        ws.cell(row, 28, "010-0000-0000")
        ws.cell(row, 29, "12345")
        ws.cell(row, 30, "서울시 어딘가")
        ws.cell(row, 31, "")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_process_pear_and_process_outputs():
    data = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1세트", "name": "정상건", "qty": 2},
        {"product": "콜라비 정품", "option": "3kg"},
        {"product": COUPANG_PRODUCT.replace("5kg", "10kg"), "option": "1세트", "name": "미등록옵션"},
    ])
    result = kolrabi_order.process_pear(data)
    assert result is not None
    _bytes, filename, stats = result
    assert "햇배선물세트" in filename
    assert stats["total"] == 1
    assert stats["product"] == "햇 배 선물세트(제주다팜)"
    assert any("미등록옵션" in s for s in stats["needs_check"])   # 조용히 사라지지 않는다
    ws = load_workbook(BytesIO(_bytes)).active
    assert ws.cell(2, 8).value == JEJU_NAME                         # 발주명 = 제주다팜 구성상품명
    assert ws.cell(2, 2).value == "정상건" and str(ws.cell(2, 9).value) == "2"

    outputs = kolrabi_order.process_outputs(data)
    labels = {st.get("product") for _b, _f, st in outputs}
    assert "햇 배 선물세트(제주다팜)" in labels and "콜라비" in " ".join(str(l) for l in labels)


# ── 운송장 의미키 ──
def test_tracking_semantic_keys_match_both_sides():
    ol = _semantic_option_keys(JEJU_NAME, "")
    dl = _semantic_option_keys(COUPANG_PRODUCT, "1세트")
    assert "pear:1:특대과:5kg" in ol and "pear:1:특대과:5kg" in dl
    # 보자기 미동봉·대과는 다른 키 (같은 사람이 두 종류 사도 섞이지 않는다)
    plain = _semantic_option_keys("햇 배 선물세트 대과 7.5kg (13~16과)", "")
    assert "pear:0:대과:7.5kg" in plain and "pear:1:특대과:5kg" not in plain


# ── 쿠팡 D열 택배사 표기 (CJ) ──
def test_cj_variants_normalize_to_coupang_name():
    for raw in ("CJ대한통운", "CJ 대한통운", "cj대한통운", "씨제이대한통운", "씨제이 대한통운",
                "대한통운", "CJ택배", "CJ대한통운택배", "CJ대한통운(주)", "CJGLS", "CJ Logistics"):
        assert normalize_courier_name(raw) == COUPANG_CJ_NAME == "CJ 대한통운", raw
        assert coupang_courier_name(raw) == "CJ 대한통운", raw
    # 다른 택배사는 건드리지 않는다
    assert coupang_courier_name("롯데") == "롯데택배"
    assert coupang_courier_name("한진택배") == "한진택배"
    assert coupang_courier_name("우체국택배") == "우체국"


def _orderlist(rows):
    """[(name, phone, address, courier, tracking, product)] → 제주다팜 회신(K이름·M전화·O주소·Q택배사·R운송장, E=발주명)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["헤더"] * 18)
    for name, phone, addr, courier, tracking, product in rows:
        row = [""] * 18
        row[10], row[12], row[14], row[16], row[17] = name, phone, addr, courier, tracking
        row[4] = product
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dl(rows):
    """[(name, phone, address, product, option)] → DeliveryList(AA·AB·AD·K·L, E 비움)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["헤더"] * 31)
    for name, phone, addr, product, option in rows:
        row = [""] * 31
        row[26], row[27], row[29], row[10], row[11] = name, phone, addr, product, option
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracking_input_writes_cj_for_pear_even_if_reply_misspells_or_omits_courier():
    ol = _orderlist([
        ("김배님", "010-1111-0001", "서울 A", "씨제이대한통운", "410000000001", JEJU_NAME),  # 표기 변형
        ("박배님", "010-1111-0002", "서울 B", "", "410000000002", JEJU_NAME),              # 택배사 빈칸
        ("이콜님", "010-1111-0003", "서울 C", "", "255200000003", "콜라비 정품 3kg"),      # 배 아님·빈칸
    ])
    dl = _dl([
        ("김배님", "01011110001", "서울A", COUPANG_PRODUCT, "1세트"),
        ("박배님", "01011110002", "서울B", COUPANG_PRODUCT, "1세트"),
        ("이콜님", "01011110003", "서울C", "콜라비 정품", "3kg 1박스"),
    ])
    out, _fn, stats = tracking_process(ol, dl)
    assert stats["filled"] == 3 and stats["skipped"] == 0, stats
    ws = load_workbook(BytesIO(out)).active
    d = [ws.cell(r, 4).value for r in range(2, 5)]
    e = [ws.cell(r, 5).value for r in range(2, 5)]
    assert e == ["410000000001", "410000000002", "255200000003"]
    assert d[0] == "CJ 대한통운"          # 씨제이대한통운 → 쿠팡 표기
    assert d[1] == "CJ 대한통운"          # 빈칸 → 햇 배 선물세트 기본 CJ
    assert d[2] in (None, "")             # 콜라비 빈칸은 종전대로 손대지 않음
