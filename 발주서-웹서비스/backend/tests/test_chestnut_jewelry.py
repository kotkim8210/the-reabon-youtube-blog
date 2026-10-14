"""햇밤(쥬얼리프룻) 발주·운송장 (2026-10-14 신규).

쥬얼리 공급처 상품 '★파격특가★ [P2CB] 햇밤' 구성상품 24종: '햇밤 {특·대·중·소·랜덤} {1·2·3·5}kg' + 등급 없는 '햇밤 {kg}kg'.
쿠팡 '국내산 햇밤 생알밤 군밤 찐밤용' 옵션 '1kg 1박스'(등급 없음) → '햇밤 1kg'. 발송 택배사는 CJ대한통운.
'밤'만으로 잡으면 미니밤호박(제주다팜)·밤고구마(해달)와 섞이므로 반드시 배타 판정.
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.processors import myeongi_order, myeongi_tracking
from app.processors.myeongi_tracking import is_jewelryfruit_tracking_target

COUPANG_PRODUCT = "국내산 햇밤 생알밤 군밤 찐밤용"   # DeliveryList(2026-10-14) 실측


def test_chestnut_routing_and_exclusions():
    assert myeongi_order.is_jewelry_chestnut_order(COUPANG_PRODUCT, "1kg 1박스") is True
    assert myeongi_order.is_jewelry_chestnut_order(COUPANG_PRODUCT, "5kg 1박스") is True
    # 미니밤호박(제주다팜)·밤고구마(해달)·다른 쥬얼리 품목은 햇밤이 아니다
    assert myeongi_order.is_jewelry_chestnut_order("제주 미니밤호박 보우짱", "1박스 로얄 정품 3kg") is False
    assert myeongi_order.is_jewelry_chestnut_order("해남 햇 밤고구마", "1박스 3kg") is False
    assert myeongi_order.is_jewelry_chestnut_order("햇 백도 딱딱이복숭아", "1박스 1kg") is False
    assert myeongi_order.is_jewelry_chestnut_order("국내산 김천 피자두 특품", "1박스 1kg") is False


def test_chestnut_vendor_option_name():
    # 등급 없는 쿠팡 옵션 → 등급 없는 공급처 옵션
    assert myeongi_order.convert_option("1kg 1박스", COUPANG_PRODUCT) == "햇밤 1kg"
    assert myeongi_order.convert_option("2kg 1박스", COUPANG_PRODUCT) == "햇밤 2kg"
    # 옵션에 등급 토큰이 있으면 '햇밤 {등급} {kg}kg'
    assert myeongi_order.convert_option("1박스 특 3kg", COUPANG_PRODUCT) == "햇밤 특 3kg"
    assert myeongi_order.convert_option("랜덤 5kg 1박스", COUPANG_PRODUCT) == "햇밤 랜덤 5kg"
    assert myeongi_order.convert_option("중 2kg", COUPANG_PRODUCT) == "햇밤 중 2kg"
    # '특품'·'대용량'처럼 수식어 속 글자는 등급이 아니다
    assert myeongi_order.convert_option("특품 1kg 1박스", COUPANG_PRODUCT) == "햇밤 1kg"
    assert myeongi_order.convert_option("대용량 3kg", COUPANG_PRODUCT) == "햇밤 3kg"
    # 상품명에만 등급 비슷한 글자가 있어도 옵션만 본다
    assert myeongi_order.convert_option("1kg 1박스", "국내산 햇밤 특대 생알밤") == "햇밤 1kg"
    # 토스/올웨이즈 경로(passthrough)도 같은 품목명
    assert myeongi_order.jewelry_passthrough_product(COUPANG_PRODUCT, "1kg 1박스", set()) == "햇밤 1kg"


def test_chestnut_is_tracking_target():
    assert is_jewelryfruit_tracking_target(COUPANG_PRODUCT, "1kg 1박스") is True
    assert is_jewelryfruit_tracking_target("제주 미니밤호박 보우짱", "1박스 로얄 정품 3kg") is False


def _delivery(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 3, r.get("order_no", f"OID{i}"))       # C 주문번호
        ws.cell(row, 11, r["product"])                       # K 상품명
        ws.cell(row, 12, r["option"])                        # L 옵션
        ws.cell(row, 23, r.get("qty", 1))                    # W 수량
        ws.cell(row, 27, r.get("name", f"수취인{i}"))         # AA 수취인
        ws.cell(row, 28, r.get("phone", "0502-5289-0408"))   # AB 전화
        ws.cell(row, 29, "12345")                            # AC 우편번호
        ws.cell(row, 30, r.get("address", "경상남도 창원시"))  # AD 주소
        ws.cell(row, 31, r.get("memo", "문 앞"))              # AE 배송메모
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_process_outputs_chestnut_rows_and_filename():
    data = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1kg 1박스", "name": "시청어린이집"},
        {"product": COUPANG_PRODUCT, "option": "1박스 특 3kg", "name": "박테스트", "qty": 2},
        {"product": "제주 미니밤호박 보우짱", "option": "1박스 로얄 정품 3kg"},   # 제주다팜 → 발주 대상 아님
    ])
    output, filename, stats = myeongi_order.process(data)

    assert "햇밤" in filename
    assert stats["total"] == 2
    assert {o["vendor_option_name"] for o in stats["options"]} == {"햇밤 1kg", "햇밤 특 3kg"}

    ws = load_workbook(BytesIO(output)).active
    rows = [(ws.cell(r, 3).value, ws.cell(r, 7).value, ws.cell(r, 8).value) for r in range(2, 4)]
    assert rows == [("시청어린이집", "햇밤 1kg", "1"), ("박테스트", "햇밤 특 3kg", "2")]
    assert ws.cell(2, 2).value == "아이티소프트"
    assert ws.cell(2, 9).value == "식품애착"


def _orderlist(rows: list[dict]) -> bytes:
    """쥬얼리 orderlist 회신: D 주문번호·G 상품명·K 수령인·L 옵션·M 연락처·O 주소·Q 택배사·R 운송장."""
    wb = Workbook()
    ws = wb.active
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row, 4, r["order_no"])
        ws.cell(row, 7, r.get("product", "햇밤 1kg"))
        ws.cell(row, 11, r["name"])
        ws.cell(row, 12, r.get("option", ""))
        ws.cell(row, 13, r.get("phone", "0502-5289-0408"))
        ws.cell(row, 15, r.get("address", "경상남도 창원시"))
        ws.cell(row, 17, r.get("courier", ""))
        ws.cell(row, 18, r["tracking"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracking_chestnut_defaults_to_cj_and_normalizes_reply():
    delivery = _delivery([
        {"product": COUPANG_PRODUCT, "option": "1kg 1박스", "name": "시청어린이집", "order_no": "13103573092164"},
        {"product": COUPANG_PRODUCT, "option": "2kg 1박스", "name": "김회신", "order_no": "13103573092165"},
        {"product": "국내산 김천 피자두 특품", "option": "1박스 1kg", "name": "이자두", "order_no": "13103573092166"},
    ])
    orderlist = _orderlist([
        {"order_no": "13103573092164", "name": "시청어린이집", "tracking": "600012345671"},            # 택배사 빈칸
        {"order_no": "13103573092165", "name": "김회신", "tracking": "600012345672", "courier": "CJ대한통운"},
        {"order_no": "13103573092166", "name": "이자두", "tracking": "400012345673", "product": "피자두 1kg"},  # 빈칸 → 롯데 기본
    ])
    output, filename, stats = myeongi_tracking.process(orderlist, delivery)
    assert stats == {"filled": 3, "skipped": 0}
    assert "햇밤" in filename and "피자두" not in filename   # 피자두는 송장 파일명 라벨 없음(기존 동작)

    ws = load_workbook(BytesIO(output)).active
    assert [(ws.cell(r, 4).value, ws.cell(r, 5).value) for r in range(2, 5)] == [
        ("CJ 대한통운", "600012345671"),   # 햇밤 + 회신 택배사 빈칸 → 쿠팡 인식 표기 CJ
        ("CJ 대한통운", "600012345672"),   # 회신 'CJ대한통운'(붙여쓰기) → 띄어쓰기 표기로 정규화
        ("롯데택배", "400012345673"),      # 다른 쥬얼리 품목은 종전 기본값 유지
    ]


def test_toss_jewelry_collect_gate_accepts_chestnut(monkeypatch):
    """토스 수집 키워드 게이트(수박·참외·복숭아)에 햇밤이 없으면 passthrough까지 못 가 조용히 0건 — 리뷰 지적(2026-10-14)."""
    import asyncio

    from app.processors import tomato_order
    from app.toss.client import toss_client

    async def fake_get_orders(**_kwargs):
        return [
            {"productName": COUPANG_PRODUCT, "optionName": "1kg 1박스", "receiverName": "토스손님",
             "receiverPhone": "0502-0000-0001", "quantity": 1, "orderProductStatus": "PAID"},
            {"productName": "국내산 김천 피자두 특품", "optionName": "1박스 1kg", "receiverName": "자두손님",
             "receiverPhone": "0502-0000-0002", "quantity": 1, "orderProductStatus": "PAID"},   # 토스 쥬얼리 게이트 밖(기존 동작)
        ]

    monkeypatch.setattr(toss_client, "get_orders", fake_get_orders)
    entries = asyncio.run(tomato_order.collect_toss_jewelry_orders("2026-10-13", "2026-10-14"))
    assert [(e["name"], e["product"]) for e in entries] == [("토스손님", "햇밤 1kg")]
