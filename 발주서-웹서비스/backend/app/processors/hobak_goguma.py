"""햇 호박고구마(호풍미) — 해달 발주서 생성 + 쿠팡(itsoft001) DeliveryList 운송장 입력.

꿀고구마와 같은 해달(한진) 거래처지만 쿠팡 주문은 itsoft001 계정으로 들어온다(2026-09-28).
앱의 쿠팡 API 키는 rjs007(꿀고구마) 것뿐이라 itsoft001 주문은 DeliveryList 업로드로 받고,
토스 주문은 같은 토스 계정 API에서 호박고구마만 골라 합친다.

⚠️ itsoft001 DeliveryList는 D열(택배사)이 '롯데택배'로 미리 채워져 나온다. 해달은 한진 발송이라
   송장번호만 넣으면 한진 송장이 롯데로 등록된다 → 운송장 입력 시 D열을 '한진택배'로 덮어쓴다.
"""

import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font

from app.config import TEMPLATE_DIR
from app.processors.goguma_order import (
    canonical_goguma_option,
    ensure_haedal_header,
    is_hobak_goguma_text,
)
from app.processors.haedal_tracking_parser import (
    detect_haedal_columns,
    find_courier_in_row,
    find_tracking_in_row,
)
from app.processors.tracking_match import (
    coupang_courier_name,
    name_counts,
    option_key_set,
    options_match,
    requires_option_guard,
)

KST = timezone(timedelta(hours=9))

SECTION = "hobak_goguma"          # 발주 이력(중복발주 방지) 섹션 키
PRODUCT_LABEL = "햇 호박고구마(해달)"
COURIER = "한진택배"               # 해달 발송 택배사 — 쿠팡 D열 표기
SENDER_NAME = "식품애착"
SENDER_PHONE = "010-5700-7756"
ORIGIN_ADDRESS = "전라남도 해남군 산이면 새상골길 "


def _norm(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def _compact(value) -> str:
    return re.sub(r"\s+", "", str(value or "").strip())


def _digits(value) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _zip5(value) -> str:
    text = _norm(value)
    if not text:
        return ""
    try:
        return str(int(float(text))).zfill(5)
    except (ValueError, TypeError):
        return text.zfill(5)


def _qty(value) -> int:
    try:
        return int(float(value)) if value not in (None, "") else 1
    except (ValueError, TypeError):
        return 1


def is_hobak_row(product_name: object, option_text: object) -> bool:
    """쿠팡 DeliveryList 행이 호박고구마인지 (K 상품명 + L 옵션)."""
    return is_hobak_goguma_text(product_name, option_text)


def hobak_vendor_name(product_name: object, option_text: object) -> str:
    """쿠팡/토스 옵션 → 해달 품목명 '호박고구마 {kg}Kg ({등급})'.

    등급·중량을 못 읽으면 원문을 정리한 문자열이 돌아온다(호출부가 needs_check로 알림).
    """
    return canonical_goguma_option(_norm(option_text), _norm(product_name))


def _is_parsed(vendor_name: str) -> bool:
    return bool(re.fullmatch(r"호박고구마 \d+Kg \(.+\)", vendor_name or ""))


def process_order(
    delivery_file_bytes: bytes | None = None,
    toss_entries: list[dict] | None = None,
) -> tuple[bytes, str, dict]:
    """해달 한진양식 발주서 생성 — 쿠팡 itsoft001 DeliveryList 호박고구마 행 + 토스 호박고구마 entries.

    toss_entries 필드는 goguma_order.collect_toss_orders와 동일(name/phone/zipcode/address/qty/product/memo/order_id).
    """
    toss_entries = list(toss_entries or [])
    coupang_rows: list[dict] = []
    needs_check: list[str] = []

    if delivery_file_bytes:
        dl_ws = load_workbook(filename=BytesIO(delivery_file_bytes), data_only=True).active
        for row in dl_ws.iter_rows(min_row=2):
            name = _norm(row[26].value) if len(row) > 26 else ""
            product_name = _norm(row[10].value) if len(row) > 10 else ""
            option = _norm(row[11].value) if len(row) > 11 else ""
            if not name or not is_hobak_row(product_name, option):
                continue
            vendor = hobak_vendor_name(product_name, option)
            if not _is_parsed(vendor):
                needs_check.append(f"{name}({option or product_name}) — 호박고구마 중량·등급을 못 읽음, 발주서 품목명 확인")
            coupang_rows.append({
                "name": name,
                "phone": _norm(row[27].value) if len(row) > 27 else "",
                "zipcode": _zip5(row[28].value) if len(row) > 28 else "",
                "address": _norm(row[29].value) if len(row) > 29 else "",
                "qty": _qty(row[22].value if len(row) > 22 else ""),
                "product": vendor,
                "memo": _norm(row[30].value) if len(row) > 30 else "",
                "order_id": _norm(row[2].value) if len(row) > 2 else "",
                "option": option,   # 중복발주 이력 키 = 주문번호|L열 원문 (filter_delivery_by_issued와 동일 출처)
            })

    for entry in toss_entries:
        if not _is_parsed(entry.get("product") or ""):
            needs_check.append(
                f"{entry.get('name') or '이름없음'}(토스 {entry.get('product') or ''}) — 호박고구마 중량·등급을 못 읽음, 발주서 품목명 확인"
            )

    template_path = TEMPLATE_DIR / "해달_발주서_한진양식.xlsx"
    if not template_path.exists():
        raise FileNotFoundError(f"템플릿 파일을 찾을 수 없습니다: {template_path}")
    tmpl_wb = load_workbook(filename=str(template_path))
    for sheet_name in tmpl_wb.sheetnames[1:]:
        del tmpl_wb[sheet_name]
    ws = tmpl_wb[tmpl_wb.sheetnames[0]]
    ensure_haedal_header(ws)

    font11 = Font(size=11)
    row_idx = 2
    for r in range(2, ws.max_row + 2):
        if ws.cell(row=r, column=1).value is None:
            row_idx = r
            break

    option_totals: dict[str, dict] = {}
    for entry in coupang_rows + toss_entries:
        qty_int = _qty(entry.get("qty"))
        mapping = {
            1: entry.get("name", ""),
            2: entry.get("phone", ""),
            5: _zip5(entry.get("zipcode")),
            6: entry.get("address", ""),
            7: SENDER_NAME,
            8: SENDER_PHONE,
            12: ORIGIN_ADDRESS,
            13: str(qty_int),
            14: entry.get("product", ""),
            16: "선불",
            19: entry.get("memo", ""),
        }
        for col, value in mapping.items():
            ws.cell(row=row_idx, column=col, value=value).font = font11
        row_idx += 1

        key = entry.get("option") or entry.get("product") or "호박고구마"
        bucket = option_totals.setdefault(
            key,
            {
                "coupang_option_keyword": key,
                "vendor_option_name": entry.get("product") or "호박고구마",
                "quantity": 0,
                "orders": [],
            },
        )
        bucket["quantity"] += qty_int
        if entry.get("order_id"):
            bucket["orders"].append({"order_id": entry["order_id"], "quantity": qty_int})

    output = BytesIO()
    tmpl_wb.save(output)
    output.seek(0)

    now = datetime.now(KST)
    filename = f"해달 발주서 한진양식_아이티소프트_호박고구마({now.strftime('%y%m%d')}).xlsx"
    stats = {
        "total": len(coupang_rows) + len(toss_entries),
        "coupang": len(coupang_rows),
        "toss": len(toss_entries),
        "product": PRODUCT_LABEL,
        "options": list(option_totals.values()),
    }
    if needs_check:
        stats["needs_check"] = needs_check
    return output.read(), filename, stats


def _parse_haedal_entries(haedal_bytes: bytes) -> list[dict]:
    ws = load_workbook(filename=BytesIO(haedal_bytes), data_only=True).active
    cols = detect_haedal_columns(ws)
    entries: list[dict] = []
    for row_idx in range(cols.start_row, ws.max_row + 1):
        name = _compact(ws.cell(row=row_idx, column=cols.name).value)
        if not name:
            continue
        tracking = find_tracking_in_row(ws, row_idx, cols.tracking)
        if not tracking:
            continue
        product = ws.cell(row=row_idx, column=cols.product).value
        skip_cols = tuple(c for c in (cols.name, cols.phone, cols.address, cols.product, cols.tracking) if c)
        entries.append({
            "name": name,
            "phone": _digits(ws.cell(row=row_idx, column=cols.phone).value),
            "address": _compact(ws.cell(row=row_idx, column=cols.address).value),
            "tracking": tracking,
            "courier": find_courier_in_row(ws, row_idx, cols.courier, skip_cols),
            "option_keys": option_key_set(product, hobak_vendor_name("", product)),
        })
    return entries


def process_tracking(haedal_bytes: bytes, delivery_bytes: bytes) -> tuple[bytes, str, dict]:
    """해달 회신 → 쿠팡 itsoft001 DeliveryList 호박고구마 행에 E(송장번호)·D(택배사) 입력.

    호박고구마 행만 채운다 — 같은 DeliveryList의 LA갈비·홍로 등 다른 상품 고객과 이름이 같아도
    해달 송장이 엉뚱한 행에 들어가지 않게 하기 위함.
    """
    entries = _parse_haedal_entries(haedal_bytes)
    if not entries:
        raise ValueError("해달 회신 파일에서 운송장번호를 찾을 수 없습니다.")

    dl_wb = load_workbook(filename=BytesIO(delivery_bytes))
    for sheet_name in dl_wb.sheetnames[1:]:
        del dl_wb[sheet_name]
    dl_ws = dl_wb[dl_wb.sheetnames[0]]

    by_name: dict[str, list[dict]] = defaultdict(list)
    by_phone: dict[str, list[dict]] = defaultdict(list)
    for entry in entries:
        by_name[entry["name"]].append(entry)
        if entry["phone"]:
            by_phone[entry["phone"]].append(entry)

    hobak_rows = [
        r for r in range(2, dl_ws.max_row + 1)
        if is_hobak_row(dl_ws.cell(r, 11).value, dl_ws.cell(r, 12).value)
    ]
    dl_name_counts = name_counts(_compact(dl_ws.cell(r, 27).value) for r in hobak_rows)

    used: set[int] = set()
    filled = 0
    already = 0
    skip_details: list[str] = []
    for r in hobak_rows:
        e_cell = dl_ws.cell(row=r, column=5)
        if _compact(e_cell.value):
            already += 1
            continue
        dl_name = _compact(dl_ws.cell(r, 27).value)
        if not dl_name:
            continue
        dl_phone = _digits(dl_ws.cell(r, 28).value)
        dl_address = _compact(dl_ws.cell(r, 30).value)
        dl_keys = option_key_set(dl_ws.cell(r, 12).value, hobak_vendor_name(dl_ws.cell(r, 11).value, dl_ws.cell(r, 12).value))

        candidates = [c for c in by_name.get(dl_name, []) if id(c) not in used]
        if candidates and requires_option_guard(dl_name, dl_name_counts, len(by_name.get(dl_name, []))):
            candidates = [c for c in candidates if options_match(c.get("option_keys"), dl_keys)]

        matched = None
        if candidates:
            matched = (
                next((c for c in candidates if c["phone"] == dl_phone and c["address"] == dl_address), None)
                or next((c for c in candidates if dl_phone and c["phone"] == dl_phone), None)
                or next((c for c in candidates if c["address"] == dl_address), None)
                or candidates[0]
            )
        elif dl_phone:
            # 이름이 바뀐 경우(CS 수취인 변경 등) 전화번호가 유일하게 맞으면 인정
            phone_hits = [c for c in by_phone.get(dl_phone, []) if id(c) not in used]
            if len(phone_hits) == 1:
                matched = phone_hits[0]

        if matched is None:
            skip_details.append(f"{dl_ws.cell(r, 27).value}(해달 회신에 없음)")
            continue

        e_cell.value = matched["tracking"]
        # itsoft001 DeliveryList는 D열이 '롯데택배'로 미리 채워져 있다 → 해달 택배사로 반드시 덮어쓴다.
        dl_ws.cell(row=r, column=4).value = coupang_courier_name(matched.get("courier"), COURIER) or COURIER
        used.add(id(matched))
        filled += 1

    output = BytesIO()
    dl_wb.save(output)
    output.seek(0)

    now = datetime.now(KST)
    filename = f"DeliveryList송장파일_호박고구마({now.strftime('%Y%m%d')}).xlsx"
    stats: dict = {
        "filled": filled,
        "skipped": len(skip_details),
        "hobak_rows": len(hobak_rows),
        "haedal_entries": len(entries),
    }
    if already:
        stats["already_filled"] = already
    if skip_details:
        stats["skipped_names"] = ", ".join(skip_details)
    if not hobak_rows:
        stats["needs_check"] = ["DeliveryList에 호박고구마 주문이 없습니다 — itsoft001 쿠팡 DeliveryList가 맞는지 확인하세요."]
    return output.read(), filename, stats
